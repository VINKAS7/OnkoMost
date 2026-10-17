"""End-to-end against IRIS (graph + FHIR) + LLM cache: the sample case replayed through the workflow,
checked against the Rakathon answer key (MDT zaver.docx)."""
import io
from datetime import date, timedelta

import docx
import pytest

from backend.graph.store import graph
from backend.pathway import governance as gov
from backend.seed import seed
from backend.workflow import service as svc
from backend.workflow.db import Case, PathwayVersion, session_scope, set_today
from scripts import seed_demo

from .conftest import PACIENT


@pytest.fixture()
def fresh(iris_available):
    other_before = graph.foreign_counts()
    seed_demo.reset()
    seed()
    yield
    assert graph.foreign_counts() == other_before, "non-Onkomost FHIR data must never change"


def _detail(case_id):
    return svc.case_detail(case_id)


def test_referral_alone_finds_missing_reports(fresh):
    case_id = seed_demo.run("referral")
    d = _detail(case_id)
    assert d["state"] == "in_triage" and d["queue"] == "new"
    missing = [t for t in d["tasks"] if t["kind"] == "missing_document"]
    assert len(missing) == 5
    assert d["next_task"]["key"] == "step:triage"
    # after acceptance the coordinator must book CT and blood tests
    svc.accept_case(case_id, "test")
    keys = {t["key"] for t in _detail(case_id)["tasks"] if t["status"] == "open"}
    assert {"step:ct", "step:labs"} <= keys
    # "collect prior reports" is represented by the 5 individual report requests, not a duplicate task
    assert "step:prior_reports" not in keys and sum(k.startswith("missing_document:") for k in keys) == 5
    assert _detail(case_id)["state"] == "waiting_for_documents"


def test_wrong_patient_document_is_blocked(fresh, tmp_path):
    case_id = seed_demo.run("referral")
    d = docx.Document(str(PACIENT / "03_RTG_plic_2026-09-03.docx"))
    for p in d.paragraphs:
        for r in p.runs:
            r.text = r.text.replace("635722/1013", "715301/1234")
    path = tmp_path / "rtg_other_patient.docx"
    d.save(str(path))
    doc_id = svc.add_document(case_id, path.name, path.read_bytes(), sync=True)
    det = _detail(case_id)
    doc = next(x for x in det["documents"] if x["id"] == doc_id)
    assert doc["status"] == "blocked"
    assert any(t["key"] == f"identity:{doc_id}" and t["status"] == "open" for t in det["tasks"])


def test_full_story_matches_answer_key(fresh):
    case_id = seed_demo.run("histology")
    d = _detail(case_id)
    assert d["pre_mdt_complete"] and d["state"] == "complete"
    flags = [f for f in d["findings"] if f["kind"] == "mismatch"]
    assert len(flags) == 2  # doc 09 dose date, doc 11 specimen date
    nodes = {n["id"]: n for n in d["pathway_graph"]["nodes"]}
    # green / red / yellow document status on the pathway
    assert nodes["ct"]["doc_status"] == "have" and nodes["labs"]["doc_status"] == "have"
    assert nodes["biopsy"]["doc_status"] == "conflict" and nodes["histology"]["doc_status"] == "conflict"
    assert nodes["pain"]["doc_status"] == "missing"  # pain specialist report comes after the MDT

    h = svc.handover(case_id)
    assert any(m["key"] == "performance_status" for m in h["missing"])  # ECOG not in the documents
    for sec in h["sections"]:
        for f in sec["fields"]:
            for item in f["items"]:
                assert item["verified"] or f["status"] == "unverified", f["key"]

    with session_scope() as s:
        set_today(s, date(2026, 10, 20))
    svc.validate_handover(case_id, "test doctor")
    report = svc.generate_mdt_report(case_id, "test")
    text = " ".join(f"{r['exam']} {r['result']}" for r in report["results"])
    assert "CA 19-9 312" in text
    assert "G2–G3" in report["diagnosis"]["name"] and report["diagnosis"]["icd10"] == "C25.0"
    assert "35 mm" in text
    com = " ".join(report["comorbidities"]).lower()
    for c in ("fibrilace síní", "hypertenze", "hypercholesterolémie"):
        assert c in com
    assert len(report["members"]) == 7

    # post-MDT reports (pain specialist 26. 10., psychologist 27. 10.) must not reach the MDT report
    with session_scope() as s:
        set_today(s, date(2026, 10, 27))
    seed_demo.upload(case_id, "12_Algeziologie_2026-10-26.docx")
    seed_demo.upload(case_id, "13_Psychologie_KOC2_2026-10-27.docx")
    report = svc.generate_mdt_report(case_id, "test")
    assert len(report["excluded_documents"]) == 2
    assert not any("Algeziolog" in r["exam"] or "Psycholog" in r["exam"] for r in report["results"])
    assert not any("tramadol" in m.lower() for m in report["medications"].split("\n"))


def test_feedback_closes_the_loop(fresh):
    case_id = seed_demo.run("closed")
    d = _detail(case_id)
    assert d["state"] == "closed" and d["treatments"] == ["chemotherapy"]
    gp = next(m for m in d["messages"] if m["kind"] == "feedback_gp")
    review = next(m for m in d["messages"] if m["kind"] == "feedback_review")
    assert gp["acked_on"] is not None and "neoadjuvantní" in gp["body"].lower() and "Chemoterapie" in gp["body"]
    assert "Revizní lékař" in review["to_name"]
    assert not [t for t in d["tasks"] if t["status"] == "open"]


def test_governance_majority_of_all_kocs(fresh):
    case_id = seed_demo.run("referral")  # running case on v1
    with session_scope() as s:
        set_today(s, date(2026, 11, 1))
        v1 = s.query(PathwayVersion).filter_by(status="active", pathway_id=seed_demo.pathway_id("pancreas")).one()
        pathway_id, v1_id = v1.pathway_id, v1.id
        g = dict(v1.graph_json)
    g = {"nodes": [*g["nodes"], {"id": "genetics", "type": "step", "label": "Genetické vyšetření (BRCA)",
                                  "phase": "diagnostics", "owner_role": "coordinator", "completion": "manual",
                                  "blocking": False}],
         "edges": [*g["edges"], {"source": "histology", "target": "genetics"}]}

    # 3 of 6 yes → can't pass
    draft = gov.create_draft(pathway_id, 2, g)
    with session_scope() as s:
        assert s.get(PathwayVersion, draft).status == "draft"
    p = gov.submit(draft, "BRCA", "test", date(2026, 11, 10), 2)
    for k in (1, 2, 3):
        gov.cast_vote(p, k, "yes")
    t = gov.cast_vote(p, 4, "no")
    t = gov.cast_vote(p, 5, "no")
    assert t["status"] == "voting"  # 3 yes + 1 not voted could still reach 4
    gov.cast_vote(p, 6, "no")
    with session_scope() as s:
        from backend.workflow.db import ChangeProposal
        assert s.get(ChangeProposal, p).status == "rejected"
        assert s.get(PathwayVersion, v1_id).status == "active"

    # 4 of 6 yes → new version active
    draft2 = gov.create_draft(pathway_id, 3, g)
    p2 = gov.submit(draft2, "BRCA v2", "test", date(2026, 11, 10), 3)
    for k in (1, 2, 3):
        t = gov.cast_vote(p2, k, "yes")
    assert t["status"] == "voting" and t["needed"] == 4
    t = gov.cast_vote(p2, 4, "yes")
    with session_scope() as s:
        assert s.get(PathwayVersion, draft2).status == "active"
        assert s.get(PathwayVersion, v1_id).status == "retired"
        assert s.get(Case, case_id).pathway_version_id == v1_id  # running case keeps v1


def test_new_cancer_type_goes_through_the_vote(fresh):
    with session_scope() as s:
        set_today(s, date(2026, 11, 1))
    with pytest.raises(gov.GovernanceError, match="already covered"):
        gov.create_pathway(2, "Plíce 2", "Lung 2", "C34.1")  # lung owns C34
    with pytest.raises(gov.GovernanceError, match="ICD-10"):
        gov.create_pathway(2, "Prostata", "Prostate", "prostate")

    draft = gov.create_pathway(2, "Karcinom prostaty", "Prostate cancer", "c61")
    with session_scope() as s:
        v = s.get(PathwayVersion, draft)
        pathway_id = v.pathway_id
        assert (v.status, v.version, v.based_on_version_id) == ("draft", 1, None)
        assert v.pathway.code == "prostate_cancer" and v.pathway.icd_codes == ["C61"]
        assert svc._pathway_for(s, "C61").pathway_id != pathway_id  # not routable before the vote
    gov.update_pathway_info(pathway_id, 2, "Karcinom prostaty", "Prostate cancer", "C61, C62")
    with pytest.raises(gov.GovernanceError):
        gov.update_pathway_info(pathway_id, 3, "Cizí", "Someone else's", "C61")

    p = gov.submit(draft, "Prostate cancer pathway", "new cancer type", date(2026, 11, 10), 2)
    with session_scope() as s:
        from backend.workflow.db import ChangeProposal
        view = gov.proposal_view(s, s.get(ChangeProposal, p))
        assert view["new_pathway"] and view["base_version"] == 0 and view["icd_codes"] == ["C61", "C62"]
        assert "feedback" in view["diff"]["added_nodes"]
    with pytest.raises(gov.GovernanceError, match="already covered"):
        gov.create_pathway(3, "Varle", "Testis", "C62")  # the proposal being voted on claims C62
    for k in (1, 2, 3, 4):
        gov.cast_vote(p, k, "yes")
    with session_scope() as s:
        assert s.get(PathwayVersion, draft).status == "active"
        assert svc._pathway_for(s, "C61").id == draft


def test_discarding_a_new_cancer_type_draft_removes_it(fresh):
    draft = gov.create_pathway(2, "", "Melanoma", "C43")
    with pytest.raises(gov.GovernanceError):
        gov.discard_draft(draft, 3)  # not the author
    gov.discard_draft(draft, 2)
    with session_scope() as s:
        from backend.workflow.db import Pathway
        assert not s.query(Pathway).filter_by(code="melanoma").first()


def test_lung_pathway_extent_decision(fresh):
    """ESMO lung pathway: the diagnostic MDT's extent decision gates the radical work-up (PET-CT, lung function,
    mediastinal staging); biopsy and brain MRI apply to both branches."""
    seed_demo.run("referral")  # reuse the patient, but on the lung pathway via a second case
    f = PACIENT / "07_PL_zadost_KOC2_2026-10-02.docx"
    lung = svc.create_referral(1, [(f.name, f.read_bytes())], sync=True, pathway_id=seed_demo.pathway_id("lung"),
                               patient_name="Alena Procházková", rc="635722/1013")
    nodes = {n["id"]: n for n in _detail(lung)["pathway_graph"]["nodes"]}
    assert nodes["pet_ct"]["doc_status"] == "conditional" and nodes["pft"]["doc_status"] == "conditional"
    svc.choose_decision(lung, "extent", ["metastatic"], "test")
    nodes = {n["id"]: n for n in _detail(lung)["pathway_graph"]["nodes"]}
    assert all(nodes[k]["status"] == "skipped" for k in ("pet_ct", "pft", "mediastinal"))
    assert nodes["biopsy"]["doc_status"] == "missing" and nodes["brain_mri"]["doc_status"] == "missing"
    with pytest.raises(svc.WorkflowError):
        svc.choose_decision(lung, "extent", ["radical", "metastatic"], "test")  # single choice only


def test_mdt_summary_for_the_sample_case(fresh):
    """The one-sheet MDT summary: cancer details, markers, open issues, charts – only documents up to the MDT date."""
    from backend.reports import summary as sm
    case_id = seed_demo.run("ready")
    s = svc.mdt_summary(case_id, "cs")
    text = sm.to_text(s)
    assert "Duktální adenokarcinom pankreatu" in text and "MKN-10: C25.0" in text and "Grade: G2–G3" in text
    assert "CA 19-9 | 312 kU/l (15. 10. 2026)" in text  # tumour markers in the lab table
    assert "Chybí údaj: Stav výkonnosti (ECOG / Karnofsky)" in text  # a clinical gap for the MDT
    assert "e-mail" not in text.split("PŘED ROZHODNUTÍM OVĚŘIT")[1].split("DŮLEŽITÁ")[0].lower()  # optional form fields
    assert "Algeziolog (volitelné) | chybí" in text  # as of the MDT date: the pain report (26. 10.) comes later
    assert sum(1 for b in s["blocks"] if b[0] == "img") == 2  # lab trends + journey timeline
    d = docx.Document(io.BytesIO(sm.to_docx(s)))
    assert len(d.inline_shapes) == 2 and any("Nádorové markery" in c.text for t in d.tables for r in t.rows for c in r.cells)
    assert "PATIENT JOURNEY" in sm.to_text(svc.mdt_summary(case_id, "en"))
    with session_scope() as s_:
        from backend.workflow.db import AuditLog
        assert s_.query(AuditLog).filter_by(case_id=case_id, action="mdt.summary_exported").count() == 2


def test_graphrag_finds_tumour_marker(fresh):
    case_id = seed_demo.run("closed")  # includes the post-MDT pain (26. 10.) and psychology (27. 10.) reports
    from backend.rag.ask import retrieve
    with session_scope() as s:
        p_uuid = s.get(Case, case_id).patient_uuid
    hits = retrieve(p_uuid, ["CA 19-9"], date(2026, 10, 21))
    assert hits and any("312" in h["text"] for h in hits)
    assert retrieve(p_uuid, ["tramadol"], date(2026, 10, 31))
    # the pain report is after the MDT cut-off: neither the keyword nor the vector leg may return it
    hits = retrieve(p_uuid, ["tramadol"], date(2026, 10, 21), question="What pain medication does she take?")
    assert hits and all(h["date"] is None or h["date"] <= date(2026, 10, 21) for h in hits)


def test_every_chunk_has_a_vector(fresh):
    from backend.iris import db
    from backend.rag import index
    if not index.embed(["test"], "passage"):
        pytest.skip("embedding model not reachable")
    seed_demo.run("ct")
    total, with_vector = db.sql(f"SELECT COUNT(*), COUNT(Embedding) FROM {index.TABLE}")[0]
    assert total and with_vector == total


def test_identity_lives_only_in_fhir(fresh):
    from sqlalchemy import inspect

    from backend.fhir.client import fhir
    from backend.graph.store import G
    from backend.iris import db
    from backend.workflow.db import engine
    case_id = seed_demo.run("referral")
    with session_scope() as s:
        case = s.get(Case, case_id)
        p_uuid = case.patient_uuid
        assert case.patient_name == "Alena Procházková" and case.rc == "635722/1013"  # read from FHIR
    assert fhir.read("Patient", p_uuid)["identifier"][0]["value"] == "635722/1013"
    assert fhir.count("DocumentReference", patient=p_uuid) >= 1
    assert "rc" not in {c["name"] for c in inspect(engine).get_columns("cases")}
    texts = [r[0] for r in db.sql("SELECT Text FROM Onk.DocChunk WHERE PatientUuid = ?", [p_uuid])]
    assert texts and not any("635722" in t or "Procházková" in t for t in texts)
    with db.native() as n:
        stored = [v for _, v in n.iterator(G, "P", p_uuid, "D")] + [v for _, v in n.iterator(G, "P", p_uuid, "F")]
    assert stored and not any("635722" in v for v in stored)


def _docx_text(path):
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        parts += [c.text for r in t.rows for c in r.cells]
    return "\n".join(parts)


def test_final_mdt_record_has_conclusion(fresh):
    case_id = seed_demo.run("mdt")  # MDT conclusion imported and decision recorded
    with session_scope() as s:
        path = svc._abs(s.get(Case, case_id).mdt_report_path)
    text = _docx_text(path)
    assert "Zápis z jednání multidisciplinárního týmu" in text
    assert "neoadjuvantní systémovou chemoterapii" in text          # 6 Závěr MDT
    assert "Neoadjuvantní chemoterapie: termín na onkologické" in text  # 7 Doporučení MDT
    assert "Rozhodnutí o léčbě: Chemoterapie" in text


def test_feedback_attaches_final_record_without_import(fresh):
    case_id = seed_demo.run("ready")
    with session_scope() as s:
        set_today(s, date(2026, 10, 21))
    svc.record_mdt_decision(case_id, ["chemotherapy", "surgery"], "Hraničně resekabilní nález, nejprve chemoterapie.",
                            ["Onkologické konzilium 4. 11. 2026"], "MDT", members=["MUDr. Lenka Marková"])
    svc.send_feedback(case_id, "test")
    gp = next(m for m in _detail(case_id)["messages"] if m["kind"] == "feedback_gp")
    record = next(a for a in gp["attachments"] if a.get("report_path"))
    text = _docx_text(svc._abs(record["report_path"]))
    assert "Hraničně resekabilní nález, nejprve chemoterapie." in text and "Chemoterapie + Chirurgická léčba" in text
    assert not any(a.get("document_id") for a in gp["attachments"])  # no signed original was imported


def test_question_to_gp_and_answer(fresh):
    case_id = seed_demo.run("accepted")
    draft = svc.draft_question(case_id)
    assert "Stav výkonnosti" in draft
    mid = svc.post_message(case_id, "coordinator", "Bc. Lucie Dvořáková", draft)
    task = next(t for t in _detail(case_id)["tasks"] if t["key"] == f"msg:{mid}")
    assert task["status"] == "open" and task["external"] and "Řezníček" in task["owner_name"]
    svc.post_message(case_id, "gp", "MUDr. Tomáš Řezníček", "ECOG 1, e-mail: alena@example.cz")
    task = next(t for t in _detail(case_id)["tasks"] if t["key"] == f"msg:{mid}")
    assert task["status"] == "done"
    # a document uploaded by the GP also answers an open question
    mid2 = svc.post_message(case_id, "coordinator", "Bc. Lucie Dvořáková", "Prosíme o RTG plic.")
    f = PACIENT / "03_RTG_plic_2026-09-03.docx"
    svc.add_document(case_id, f.name, f.read_bytes(), source="gp", actor="MUDr. Tomáš Řezníček", sync=True)
    assert next(t for t in _detail(case_id)["tasks"] if t["key"] == f"msg:{mid2}")["status"] == "done"


def test_review_doctor_question(fresh):
    case_id = seed_demo.run("feedback")
    review = next(m for m in _detail(case_id)["messages"] if m["kind"] == "feedback_review")
    mid = svc.review_question(review["ack_token"], "MUDr. Revizní", "Prosím o kopii histologie.")
    task = next(t for t in _detail(case_id)["tasks"] if t["key"] == f"msg:{mid}")
    assert task["status"] == "open" and not task["external"] and task["priority"] == "high"
    assert any(n["level"] == "question" and n["case_id"] == case_id for n in svc.notifications(2))
    svc.post_message(case_id, "coordinator", "Bc. Lucie Dvořáková", "Histologie v příloze.", to_role="review")
    assert next(t for t in _detail(case_id)["tasks"] if t["key"] == f"msg:{mid}")["status"] == "done"
    # the review doctor's receipt does not close the GP feedback step
    assert next(n for n in _detail(case_id)["pathway_graph"]["nodes"] if n["id"] == "feedback")["status"] != "done"


def test_notifications_list_overdue(fresh):
    case_id = seed_demo.run("histology")  # demo date 19. 10.; supportive steps were due 17. 10.
    items = svc.notifications(2)
    assert any(n["level"] == "overdue" and n["case_id"] == case_id for n in items)
    assert not svc.notifications(1)  # nothing for another centre

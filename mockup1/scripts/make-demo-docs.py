"""Generate the demo outpatient report for the GP document scan (data/demo/ambulantni-zprava-novak.pdf + .docx).

The report documents exactly what Josef Novák's referral is still missing on the GP screen (ECOG, CA 19-9, blood
count + coagulation, medication without anticoagulants) and, as traps for the LLM, says that CT has not been done
and histology is not available yet. Fictional patient, facility and doctors.

    python3 scripts/make-demo-docs.py            # from mockup1/; needs google-chrome (PDF) and python-docx (DOCX)
"""
from __future__ import annotations

import html
import subprocess
import tempfile
from pathlib import Path

OUT = Path(__file__).resolve().parent.parent / "data" / "demo"
NAME = "ambulantni-zprava-novak"

HEAD = [
    ("Pacient", "Josef Novák, nar. 5. 1. 1966, RČ 660105/••••, pojišťovna 111"),
    ("Datum vyšetření", "6. 1. 2026"),
    ("Odesílající", "MUDr. Jana Dvořáková, praktický lékař, Benešov"),
]
SECTIONS = [
    ("Důvod vyšetření", ["Nebolestivý ikterus od 5. 1. 2026, odeslán praktickým lékařem k došetření."]),
    ("Anamnéza", [
        "DM 2. typu na PAD, arteriální hypertenze. Úbytek hmotnosti 4 kg za 2 měsíce, acholická stolice, tmavá moč. "
        "Nekuřák, alkohol příležitostně. RA: otec karcinom prostaty, jinak bez nádorů GIT.",
    ]),
    ("Farmakologická anamnéza", [
        "Metformin 1000 mg 1-0-1, Ramipril 5 mg 1-0-0. Antikoagulancia ani antiagregancia neužívá. Alergie: penicilin.",
    ]),
    ("Objektivně", [
        "Ikterus kůže a sklér, břicho měkké, nebolestivé, bez hmatné rezistence, játra nezvětšena.",
        "Celkový stav: soběstačný, chodí bez pomoci, běžné denní činnosti zvládá, těžší fyzickou práci již ne, "
        "většinu dne mimo lůžko – výkonnostní stav ECOG 1.",
    ]),
]
LABS_TITLE = "Laboratoř (odběr 6. 1. 2026, OKB Nemocnice Benešov)"
LABS = [
    ("Bilirubin celkový", "186 µmol/l", "2–17", "↑"),
    ("ALP", "4,2 µkat/l", "0,7–2,2", "↑"),
    ("GMT", "7,8 µkat/l", "0,1–1,2", "↑"),
    ("ALT / AST", "3,1 / 2,4 µkat/l", "do 0,8", "↑"),
    ("CRP", "12 mg/l", "do 5", "↑"),
    ("Leukocyty", "7,9 × 10⁹/l", "4,0–10,0", ""),
    ("Hemoglobin", "128 g/l", "135–175", "↓"),
    ("Trombocyty", "245 × 10⁹/l", "150–400", ""),
    ("INR", "1,1", "0,8–1,2", ""),
    ("aPTT", "31 s", "25–38", ""),
    ("CA 19-9", "389 kU/l", "0–37", "↑"),
    ("CEA", "4,1 µg/l", "0–5", ""),
]
TAIL = [
    ("Zobrazovací vyšetření", [
        "UZ břicha 6. 1. 2026: hypoechogenní ložisko v hlavě pankreatu cca 32 mm, dilatace intra- i extrahepatálních žlučových cest.",
        "CT břicha zatím neprovedeno – doporučeno CT podle pankreatického protokolu v KOC.",
    ]),
    ("Histologie", ["Dosud není k dispozici."]),
    ("Závěr a doporučení", [
        "Obstrukční ikterus, suspektní tumor hlavy pankreatu. Doporučeno odeslání do KOC VFN Praha k CT (pankreatický "
        "protokol), ERCP s drenáží žlučových cest a EUS-FNB dle rozhodnutí KOC. Pacient poučen.",
    ]),
]
FACILITY = "Nemocnice Benešov · Interní ambulance"
SIGN = "MUDr. Pavel Kratochvíl, internista · 6. 1. 2026"
TITLE = "Ambulantní zpráva"


def make_html() -> str:
    e = html.escape
    sec = lambda t, ps: f"<h2>{e(t)}</h2>" + "".join(f"<p>{e(p)}</p>" for p in ps)
    rows = "".join(f"<tr><td>{e(a)}</td><td>{e(b)}</td><td>{e(c)}</td><td>{e(d)}</td></tr>" for a, b, c, d in LABS)
    return f"""<!doctype html><html lang="cs"><head><meta charset="utf-8"><title>{e(TITLE)}</title><style>
@page {{ size: A4; margin: 18mm 18mm 16mm; }}
body {{ font: 10.5pt/1.45 'DejaVu Sans', sans-serif; color: #111; }}
.fac {{ font-size: 9pt; color: #444; border-bottom: 1.5px solid #333; padding-bottom: 4px; }}
h1 {{ font-size: 15pt; margin: 10px 0 6px; }} h2 {{ font-size: 11pt; margin: 12px 0 3px; }} p {{ margin: 0 0 4px; }}
dl {{ display: grid; grid-template-columns: max-content 1fr; gap: 1px 12px; margin: 0 0 4px; }} dt {{ color: #444; }} dd {{ margin: 0; }}
table {{ border-collapse: collapse; width: 100%; font-size: 9.5pt; }} th, td {{ border-bottom: 1px solid #ccc; padding: 2px 6px; text-align: left; }}
th {{ background: #eee; }} .sign {{ margin-top: 18px; }}
</style></head><body>
<div class="fac">{e(FACILITY)}</div><h1>{e(TITLE)}</h1>
<dl>{''.join(f'<dt>{e(a)}:</dt><dd>{e(b)}</dd>' for a, b in HEAD)}</dl>
{''.join(sec(t, ps) for t, ps in SECTIONS)}
<h2>{e(LABS_TITLE)}</h2><table><tr><th>Vyšetření</th><th>Výsledek</th><th>Ref. meze</th><th></th></tr>{rows}</table>
{''.join(sec(t, ps) for t, ps in TAIL)}
<p class="sign">{e(SIGN)}</p></body></html>"""


def make_pdf(path: Path) -> None:
    with tempfile.TemporaryDirectory() as d:
        src = Path(d) / "report.html"
        src.write_text(make_html(), encoding="utf-8")
        subprocess.run(["google-chrome", "--headless=new", "--disable-gpu", "--no-sandbox", "--no-pdf-header-footer",
                        f"--print-to-pdf={path}", src.as_uri()], check=True, capture_output=True, timeout=120)


def make_docx(path: Path) -> None:
    import docx
    doc = docx.Document()
    doc.add_paragraph(FACILITY)
    doc.add_heading(TITLE, level=1)
    for a, b in HEAD:
        doc.add_paragraph(f"{a}: {b}")
    for t, ps in SECTIONS:
        doc.add_heading(t, level=2)
        for p in ps:
            doc.add_paragraph(p)
    doc.add_heading(LABS_TITLE, level=2)
    table = doc.add_table(rows=1, cols=4)
    for cell, h in zip(table.rows[0].cells, ("Vyšetření", "Výsledek", "Ref. meze", "")):
        cell.text = h
    for row in LABS:
        for cell, v in zip(table.add_row().cells, row):
            cell.text = v
    for t, ps in TAIL:
        doc.add_heading(t, level=2)
        for p in ps:
            doc.add_paragraph(p)
    doc.add_paragraph(SIGN)
    doc.save(path)


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    make_pdf(OUT / f"{NAME}.pdf")
    make_docx(OUT / f"{NAME}.docx")
    print(f"wrote {OUT / NAME}.pdf and .docx")

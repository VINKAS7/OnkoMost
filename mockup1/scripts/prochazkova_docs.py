"""Demo documents for the Procházková case (Rakathon 2026 competition data, fictional patient).

    handover-prochazkova.pdf/.docx   structured handover per Handover_obsah.docx, as of the MDT on 21. 10. 2026:
                                     built from reports 01–11 only (12–13 were written after the board), missing data
                                     marked "NEUVEDENO", contradictions between reports marked "NESOULAD"
    mdt-zapis-prochazkova.pdf/.docx  MDT record per MDT_zapis.docx, filled from the sample conclusion (MDT zaver.docx)

The section texts come from data/handover-prochazkova-mdt.json and data/mdt-structured-prochazkova.json, so the
uploaded document and the prototype's pre-baked state say the same. Called by make-demo-docs.py (prochazkova).
"""
from __future__ import annotations

import html
import json
import subprocess
import tempfile
from pathlib import Path

DATA = Path(__file__).resolve().parent.parent / "data"
FICTIONAL = "FIKTIVNÍ PACIENTKA – testovací data pro Rakathon 2026"
FLAG = {"missing": "NEUVEDENO", "conflict": "NESOULAD", "note": "POZNÁMKA"}

# Provedená vyšetření: date | place | type | full result text (Handover_obsah asks for the full wording).
EXAMS = [
    ("20. 7. 2026", "Kardiologická ambulance Jesenice (MUDr. Petr Malý)", "Kardiologická kontrola, EKG, orientační echo",
     "Stabilní stav, sinusový rytmus, bez recidivy fibrilace síní. TK 132/80, TF 68/min. LK nedilatovaná, normální kinetika, "
     "EF LK 60 %, LS 42 mm, chlopně bez významné vady. Hypertenze kompenzovaná. Před plánovaným invazivním výkonem s vysokým "
     "rizikem krvácení vysadit apixaban 48 hodin předem, bez nutnosti překlenovací léčby."),
    ("3. 9. 2026", "Radiologické centrum Modřany", "RTG hrudníku PA",
     "Plicní parenchym bez ložiskových a infiltrativních změn, hily nerozšířené, bránice hladká, kostofrenické úhly volné, "
     "srdeční stín nezvětšen. Závěr: bez známek zánětlivých změn plic."),
    ("15. 9. 2026", "Poliklinika Jižní, gastroenterologie (MUDr. Ivana Křížová)", "Gastroenterologické vyšetření",
     "Tupé bolesti v epigastriu s propagací do zad, horší po jídle a v noci, nechutenství, úbytek hmotnosti 4 kg za měsíc "
     "(z 69 na 65 kg), tmavá moč, občas světlejší stolice. Subikterus sklér, břicho citlivé v epigastriu. Závěr: cholestatický "
     "syndrom; DD choledocholitiáza, stenóza žlučových cest, tumor pankreatu. Doporučeno UZ, dle výsledku CT (pankreatický "
     "protokol), případně EUS."),
    ("29. 9. 2026", "Poliklinika Jižní, RDG pracoviště", "UZ břicha",
     "Intrahepatální žlučovody lehce dilatovány, DHC 10 mm, žlučník mírně distendovaný bez konkrementů. V hlavě pankreatu "
     "hypoechogenní nehomogenní neostře ohraničené ložisko 34 × 30 mm, ductus Wirsungi 4 mm. Vztah k portomezenterickému "
     "soutoku nelze UZ posoudit. Bez volné tekutiny. Závěr: ložisková léze hlavy pankreatu suspektní z malignity."),
    ("8. 10. 2026", "KOC Nemocnice Severní, Radiologická klinika", "CT břicha a hrudníku s k. l., pankreatický protokol",
     "V hlavě pankreatu hypodenzní ložisko 35 × 31 mm. V. mesenterica superior obkružena asi 200° obvodu se zúžením lumen "
     "v délce 18 mm, proximální i distální segment vhodný k rekonstrukci; v. portae průchodná. Kontakt s a. mesenterica superior "
     "do 150° bez zúžení; truncus coeliacus a a. hepatica communis bez kontaktu; a. gastroduodenalis v kontaktu. DHC 11 mm. "
     "Peripankreatická uzlina 12 × 9 mm suspektní, paraaortální uzliny nezvětšené. Játra bez ložisek, bez ascitu. Hrudník bez "
     "ložiskových změn. Závěr: tumor hlavy pankreatu 35 mm s venózním postižením VMS a kontaktem s AMS, suspektní uzlina, bez "
     "vzdálených metastáz."),
    ("12. 10. 2026", "KOC Nemocnice Lipová, endoskopie (MUDr. Martin Svoboda)", "EUS-FNB ložiska hlavy pankreatu",
     "Hypoechogenní nehomogenní ložisko hlavy pankreatu 33 × 29 mm, 3 průchody jehlou FNB 22G, dostatečné válečky tkáně. "
     "Bez komplikací. Příprava: „apixaban vysazen 48 h před výkonem (poslední dávka 18. 10. 2026 večer)“. Doporučení: apixaban "
     "znovu od 21. 10. 2026 ráno."),
]
LABS = [
    ("Leukocyty", "12,8 ↑", "8,9", "10⁹/l"), ("Hemoglobin", "124", "118 ↓", "g/l"), ("Trombocyty", "298", "312", "10⁹/l"),
    ("Bilirubin celkový", "38 ↑", "23 ↑", "µmol/l"), ("ALT", "1,92 ↑", "1,62 ↑", "µkat/l"), ("AST", "1,45 ↑", "1,20 ↑", "µkat/l"),
    ("GGT", "4,20 ↑", "5,00 ↑", "µkat/l"), ("ALP", "3,10 ↑", "3,40 ↑", "µkat/l"), ("Albumin", "–", "36", "g/l"),
    ("CRP", "86 ↑", "18 ↑", "mg/l"), ("Kreatinin", "78", "74", "µmol/l"), ("CA 19-9", "neprovedeno", "312 ↑", "kU/l"),
    ("CEA", "neprovedeno", "4,1", "µg/l"), ("CA 125", "neprovedeno", "48 ↑", "kU/l"),
]
LABS_NOTE = ("Moč 1. 9.: bilirubin ++, leukocyty ++, sediment leukocyty 65/µl; kultivace E. coli 10⁴ CFU/ml (signifikance "
             "hraniční). Zdroje: 02 Laboratoř LabJih 1. 9. 2026, 10 Laboratoř KOC Nemocnice Lipová 15. 10. 2026.")

# A document is a list of blocks: (kind, payload).


def handover_blocks() -> tuple[str, list]:
    h = json.loads((DATA / "handover-prochazkova-mdt.json").read_text(encoding="utf-8"))
    sec = {s["id"]: s for s in h["sections"]}
    b: list = [
        ("fac", "KOC Nemocnice Lipová · koordinace KOC · Nemocniční 15, 140 00 Praha 4"),
        ("small", FICTIONAL),
        ("h1", "Strukturovaný handover – podklad pro MDT"),
        ("kv", [
            ("MDT", "MDT pro nádory GIT, 21. 10. 2026"),
            ("Stav dokumentace k", "20. 10. 2026 · zprávy 01–11 (zprávy 12 algeziologie 26. 10. a 13 psychologie 27. 10. vznikly po MDT – nezahrnuty)"),
            ("Referující lékař", "MUDr. Tomáš Řezníček, praktický lékař, Na Návsi 3, 252 41 Dolní Břežany, tel. +420 311 000 287"),
            ("Vypracovala", "Bc. Lucie Dvořáková, koordinátorka KOC (AI návrh, ověřeno koordinátorkou)"),
        ]),
        ("small", "Údaje převzaty bez domýšlení. NEUVEDENO = údaj v dokumentaci chybí, NESOULAD = zprávy si odporují – ověřit."),
        ("h2", "1  Nacionále pacienta"),
        ("kv", [("Jméno", "Alena"), ("Příjmení", "Procházková"), ("Věk", "63 let (nar. 22. 7. 1963), RČ 635722/1013"),
                ("Pohlaví", "žena"), ("Zdravotní pojišťovna", "211")]),
        ("h2", "2  Kontakty"),
        ("kv", [("Pacientka", "tel. +420 777 000 541 · Za Školou 17, 252 41 Dolní Břežany"),
                ("E-mail", "NEUVEDENO"),
                ("Kontaktní osoba", "Petra Horká (dcera), tel. +420 777 000 542 – pacientka souhlasí se sdílením informací")]),
    ]
    order = [("diagnosis", "3  Diagnóza"), ("epicrisis", "4  Epikríza (stručný souhrn nemoci)"),
             ("mdt-question", "5  Požadavek / otázka na MDT"), ("comorbidities", "6  Komorbidity (OA)"),
             ("medication", "7  Léková anamnéza (FA)"), ("allergies", "8  Alergie (AA)"),
             ("performance", "9  Stav výkonnosti (PS / Karnofského skóre)"), ("body", "10  Váha a výška"),
             ("symptoms", "11  Symptomy nemoci a jejich počátek")]
    for sid, title in order:
        s = sec[sid]
        b += [("h2", title), ("p", s["text"])]
        b += [("flag", f"{FLAG[f['kind']]}: {f['text']}") for f in s.get("flags", [])]
        b.append(("src", f"Zdroj: {s['source']}"))
    b += [("h2", "12  Provedená vyšetření"),
          ("table", (("Datum", "Místo provedení", "Typ vyšetření", "Popis výsledků (plné znění)"), EXAMS))]
    b += [("flag", f"{FLAG[f['kind']]}: {f['text']}") for f in sec["medication"].get("flags", [])]
    b += [("h2", "13  Laboratorní výsledky včetně onkomarkerů"),
          ("table", (("Vyšetření", "1. 9. 2026 (LabJih)", "15. 10. 2026 (KOC 2)", "Jednotky"), LABS)), ("p", LABS_NOTE)]
    for sid, title in [("histology", "14  Histologie"), ("other", "15  Jiná důležitá sdělení včetně přání pacienta")]:
        s = sec[sid]
        b += [("h2", title), ("p", s["text"])]
        b += [("flag", f"{FLAG[f['kind']]}: {f['text']}") for f in s.get("flags", [])]
        b.append(("src", f"Zdroj: {s['source']}"))
    b.append(("sign", "Bc. Lucie Dvořáková, koordinátorka KOC Nemocnice Lipová · 20. 10. 2026"))
    return "Strukturovaný handover – Alena Procházková", b


def mdt_blocks() -> tuple[str, list]:
    m = json.loads((DATA / "mdt-structured-prochazkova.json").read_text(encoding="utf-8"))
    r = m["record"]
    results = [
        ("29. 9. 2026", "UZ břicha", "Hypoechogenní ložisko hlavy pankreatu 34 × 30 mm; mírná dilatace žlučových cest, bez jaterních ložisek."),
        ("8. 10. 2026", "CT hrudníku a břicha, pankreatický protokol (KOC Nemocnice Severní)",
         "Tumor hlavy pankreatu 35 × 31 mm, kontakt s VMS asi 200° se zúžením lumen v délce 18 mm, kontakt s AMS do 150° bez stenózy; "
         "suspektní peripankreatická uzlina 12 × 9 mm, bez vzdálených metastáz."),
        ("12. 10. 2026", "EUS-FNB", "Ložisko hlavy pankreatu 33 × 29 mm, provedena biopsie, bez komplikací."),
        ("15. 10. 2026", "Laboratoř", "Bilirubin 23 µmol/l, ALT 1,62, GGT 5,00, ALP 3,40 µkat/l, CA 19-9 312 kU/l, albumin 36 g/l, "
         "Hb 118 g/l, CRP 18 mg/l, kreatinin 74 µmol/l."),
        ("19. 10. 2026", "Histologie", "Duktální adenokarcinom pankreatu G2–G3, perineurální invaze; CK7+, CK19+, CK20−, CDX2−."),
    ]
    b: list = [
        ("fac", f"Komplexní onkologické centrum Nemocnice Lipová · {r['board']}"),
        ("small", FICTIONAL),
        ("h1", "Zápis k jednání multidisciplinárního týmu (MDT)"),
        ("h2", "1  Základní údaje"),
        ("kv", [("Datum MDT", "21. 10. 2026"),
                ("Pacient", "Alena Procházková, RČ 635722/1013"),
                ("Pojišťovna", "211"),
                ("Telefon, e-mail", "+420 777 000 541 · e-mail NEUVEDENO"),
                ("Adresa", "Za Školou 17, 252 41 Dolní Břežany"),
                ("Váha, výška", "65 kg (15. 9. 2026), 165 cm"),
                ("Referující lékař", "MUDr. Tomáš Řezníček, praktický lékař"),
                ("Telefon, e-mail", "+420 311 000 287 · e-mail NEUVEDENO"),
                ("Případ prezentoval", r["presenter"])]),
        ("h2", "2  Členové MDT"), ("list", r["members"]),
        ("h2", "3  Diagnóza"),
        ("kv", [("Název diagnózy", "Duktální adenokarcinom hlavy pankreatu, G2–G3, s perineurální invazí; hraničně resekabilní, cM0"),
                ("Kód diagnózy (MKN-10)", r["icd10"])]),
        ("h2", "4  Komorbidity"),
        ("p", "Paroxysmální fibrilace síní na apixabanu, arteriální hypertenze, hypercholesterolémie; kardiologicky v červenci 2026 "
              "stabilní, EF LK 60 %. Výkonnostní stav (ECOG) v podkladech NEUVEDEN."),
        ("h2", "5  Výsledky stávajících vyšetření"), ("table", (("Datum", "Vyšetření", "Výsledek"), results)),
        ("h2", "6  Závěr MDT"), ("p", m["mdt"][0]["conclusion"]),
        ("h2", "7  Doporučení MDT"), ("olist", r["recommendations"]),
        ("sign", f"V Praze dne 21. 10. 2026 · {r['chair']}, předsedající MDT · Kopie: {r['copyTo']}"),
    ]
    return "Zápis z MDT – Alena Procházková", b


def to_html(title: str, blocks: list) -> str:
    e = html.escape
    out = []
    for kind, v in blocks:
        if kind == "fac": out.append(f'<div class="fac">{e(v)}</div>')
        elif kind == "small": out.append(f'<p class="small">{e(v)}</p>')
        elif kind == "h1": out.append(f"<h1>{e(v)}</h1>")
        elif kind == "h2": out.append(f"<h2>{e(v)}</h2>")
        elif kind == "p": out.append("".join(f"<p>{e(x)}</p>" for x in v.split("\n")))
        elif kind == "flag": out.append(f'<p class="flag">{e(v)}</p>')
        elif kind == "src": out.append(f'<p class="src">{e(v)}</p>')
        elif kind == "kv": out.append("<dl>" + "".join(f"<dt>{e(a)}:</dt><dd>{e(b)}</dd>" for a, b in v) + "</dl>")
        elif kind == "list": out.append("<ul>" + "".join(f"<li>{e(x)}</li>" for x in v) + "</ul>")
        elif kind == "olist": out.append("<ol>" + "".join(f"<li>{e(x)}</li>" for x in v) + "</ol>")
        elif kind == "table":
            head, rows = v
            out.append("<table><tr>" + "".join(f"<th>{e(h)}</th>" for h in head) + "</tr>"
                       + "".join("<tr>" + "".join(f"<td>{e(c)}</td>" for c in r) + "</tr>" for r in rows) + "</table>")
        elif kind == "sign": out.append(f'<p class="sign">{e(v)}</p>')
    return f"""<!doctype html><html lang="cs"><head><meta charset="utf-8"><title>{e(title)}</title><style>
@page {{ size: A4; margin: 16mm 16mm 14mm; }}
body {{ font: 10pt/1.42 'DejaVu Sans', sans-serif; color: #111; }}
.fac {{ font-size: 9pt; color: #444; border-bottom: 1.5px solid #333; padding-bottom: 4px; }}
.small {{ font-size: 8.5pt; color: #555; margin: 3px 0; }}
h1 {{ font-size: 14pt; margin: 8px 0 6px; }} h2 {{ font-size: 10.5pt; margin: 11px 0 3px; }} p {{ margin: 0 0 3px; }}
dl {{ display: grid; grid-template-columns: max-content 1fr; gap: 1px 12px; margin: 0 0 4px; }} dt {{ color: #444; }} dd {{ margin: 0; }}
table {{ border-collapse: collapse; width: 100%; font-size: 8.8pt; margin: 2px 0 4px; }}
th, td {{ border-bottom: 1px solid #ccc; padding: 2px 5px; text-align: left; vertical-align: top; }} th {{ background: #eee; }}
.flag {{ font-weight: bold; border-left: 3px solid #333; padding-left: 6px; }}
.src {{ font-size: 8.5pt; color: #555; }} .sign {{ margin-top: 14px; }} ul, ol {{ margin: 0 0 4px; padding-left: 20px; }}
</style></head><body>{''.join(out)}</body></html>"""


def to_pdf(title: str, blocks: list, path: Path) -> None:
    with tempfile.TemporaryDirectory() as d:
        src = Path(d) / "doc.html"
        src.write_text(to_html(title, blocks), encoding="utf-8")
        subprocess.run(["google-chrome", "--headless=new", "--disable-gpu", "--no-sandbox", "--no-pdf-header-footer",
                        f"--print-to-pdf={path}", src.as_uri()], check=True, capture_output=True, timeout=120)


def to_docx(title: str, blocks: list, path: Path) -> None:
    import docx
    doc = docx.Document()
    doc.core_properties.title = title
    for kind, v in blocks:
        if kind in ("fac", "small", "src", "sign"):
            doc.add_paragraph(v)
        elif kind == "h1": doc.add_heading(v, level=1)
        elif kind == "h2": doc.add_heading(v, level=2)
        elif kind == "p":
            for x in v.split("\n"):
                doc.add_paragraph(x)
        elif kind == "flag": doc.add_paragraph().add_run(v).bold = True
        elif kind == "kv":
            for a, b in v:
                doc.add_paragraph(f"{a}: {b}")
        elif kind == "list":
            for x in v:
                doc.add_paragraph(x, style="List Bullet")
        elif kind == "olist":
            for x in v:
                doc.add_paragraph(x, style="List Number")
        elif kind == "table":
            head, rows = v
            t = doc.add_table(rows=1, cols=len(head))
            t.style = "Table Grid"
            for cell, h in zip(t.rows[0].cells, head):
                cell.text = h
            for r in rows:
                for cell, c in zip(t.add_row().cells, r):
                    cell.text = c
    doc.save(path)


def build(out: Path) -> None:
    for name, make in (("handover-prochazkova", handover_blocks), ("mdt-zapis-prochazkova", mdt_blocks)):
        title, blocks = make()
        to_pdf(title, blocks, out / f"{name}.pdf")
        to_docx(title, blocks, out / f"{name}.docx")
        print(f"wrote {out / name}.pdf and .docx")

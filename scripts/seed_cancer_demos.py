"""Add two fictional MDT-ready patients – lung and breast cancer – next to the sample case, so the MDT
panel (and its body map) can be shown for every cancer type.

    python -m scripts.seed_cancer_demos            # add the patients that are not there yet
    python -m scripts.seed_cancer_demos --bodymap  # … and run the LLM body map for them

Each patient gets short synthetic Czech reports (.docx, same layout as the Rakathon sample), is referred through the
normal workflow (identity check, LLM extraction, pathway conformance), accepted, walked through the manual
diagnostic steps and given an MDT date. `python -m scripts.seed_demo` still resets everything.
"""
from __future__ import annotations

import argparse
import io
import sys
from datetime import date

import docx
from sqlalchemy import select

from backend.ingest.identity import patient_uuid
from backend.seed import seed
from backend.workflow import service as svc
from backend.workflow.db import Case, Pathway, session_scope

KOC = 2
MDT_DATE = date(2026, 10, 28)
COORD, GP = "Bc. Lucie Dvořáková", "MUDr. Tomáš Řezníček"
PRACTICE = "Ordinace praktického lékaře, Dolní Břežany"


def report(facility: list[str], title: str, patient: str, rc: str, date_line: str, sections: dict[str, list],
           author: tuple[str, str], who: str = "Pacient") -> bytes:
    """A report in the sample's layout. A section is a list of paragraphs; a list of rows becomes a table."""
    d = docx.Document()
    d.add_paragraph(f"FIKTIVNÍ {who.upper()} – ukázková data Onkomost")
    for line in facility:
        d.add_paragraph(line)
    d.add_paragraph(title)
    d.add_paragraph(f"{who}: {patient}, RČ {rc}, ZP 111")
    d.add_paragraph(f"Indikující lékař: {GP}, praktický lékař")
    d.add_paragraph(date_line)
    for head, body in sections.items():
        d.add_paragraph(head)
        for item in body:
            if isinstance(item, list):
                t = d.add_table(rows=0, cols=len(item[0]))
                for row in item:
                    cells = t.add_row().cells
                    for c, v in zip(cells, row):
                        c.text = v
            else:
                d.add_paragraph(item)
    d.add_paragraph(f"Datum: {date_line.split(': ')[-1]}")
    d.add_paragraph(author[0])
    d.add_paragraph(author[1])
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


GP_ADDR = [PRACTICE, "Náměstí 5, 252 41 Dolní Břežany", "IČP 12345678"]
KOC_RAD = ["KOC Nemocnice Lipová (fiktivní)", "Radiologická klinika", "Nemocniční 15, 140 00 Praha 4"]
KOC_PAT = ["KOC Nemocnice Lipová (fiktivní)", "Ústav patologie", "Nemocniční 15, 140 00 Praha 4"]
KOC_ENDO = ["KOC Nemocnice Lipová (fiktivní)", "Endoskopické pracoviště", "Nemocniční 15, 140 00 Praha 4"]
KOC_LAB = ["KOC Nemocnice Lipová (fiktivní)", "Oddělení klinické biochemie a hematologie", "Nemocniční 15, 140 00 Praha 4"]
LAB_HEAD = ["Vyšetření", "Výsledek", "Jednotky", "Ref. meze"]


def referral(p: dict, diag: str, story: list[str], exams: str) -> tuple[str, bytes]:
    return "01_PL_zadost.docx", report(GP_ADDR, "Epikríza a žádost o převzetí do péče KOC", p["name"], p["rc"],
                                       f"Datum: {p['referral']}", {
                                           "NYNÍ": story,
                                           "DOSAVADNÍ VYŠETŘENÍ": [exams],
                                           "ŽÁDOST": [f"Suspektní diagnóza: {diag}",
                                                      "Žádám o převzetí do péče KOC k dokončení diagnostiky a "
                                                      "stanovení dalšího postupu v rámci MDT."]},
                                       (GP, "praktický lékař"), p["who"])


PATIENTS = [
    {
        "code": "lung", "name": "Jan Dvořák", "rc": "620314/1234", "who": "Pacient", "referral": "1. 10. 2026",
        "docs": lambda p: [
            referral(p, "C34.1", ["Kuřák 40 balíčkoroků, 3 měsíce kašel, hemoptýza, úbytek 5 kg."],
                     "CT hrudníku a břicha 29. 9. 2026 (KOC Nemocnice Lipová) – tumor pravého horního laloku."),
            ("02_CT_hrudniku.docx", report(KOC_RAD, "Výsledek vyšetření – CT hrudníku a břicha s k. l.", p["name"],
                                           p["rc"], "Datum vyšetření: 29. 9. 2026", {
                "POPIS": ["V pravém horním laloku plíce spikulované ložisko 36 × 30 mm, šíře k bronchu B1. "
                          "Hilová uzlina vpravo (10R) 14 mm, suspektní. Mediastinální uzliny nezvětšeny. "
                          "Pohrudnice bez fluidothoraxu. Játra bez ložiskových změn. Nadledviny nezvětšené, bez ložisek."],
                "ZÁVĚR": ["Tumor pravého horního laloku plíce 36 mm, suspektní hilová uzlina vpravo. "
                          "Bez známek vzdálených metastáz."]}, ("MUDr. Petra Malá", "radiolog"))),
            ("03_PET_CT.docx", report(KOC_RAD, "Výsledek vyšetření – PET/CT celotělové", p["name"], p["rc"],
                                      "Datum vyšetření: 8. 10. 2026", {
                "POPIS": ["Hypermetabolický tumor pravého horního laloku plíce, SUVmax 12,4. "
                          "Zvýšená akumulace FDG v hilové uzlině vpravo (10R), SUVmax 6,1. "
                          "Nadledviny bez akumulace FDG. Játra bez ložisek zvýšené akumulace. Skelet bez ložisek."],
                "ZÁVĚR": ["Metabolicky aktivní tumor pravé plíce s postižením hilové uzliny vpravo. "
                          "Bez známek vzdálených metastáz."]}, ("MUDr. Karel Veselý", "nukleární medicína"))),
            ("04_Bronchoskopie_EBUS.docx", report(KOC_ENDO, "Protokol o výkonu – bronchoskopie a EBUS-TBNA",
                                                  p["name"], p["rc"], "Datum výkonu: 9. 10. 2026", {
                "POPIS": ["Ve vstupu do bronchu B1 vpravo infiltrace sliznice, odebrány biopsie. "
                          "EBUS-TBNA: punkce uzliny 10R, 4R a 7."],
                "ZÁVĚR": ["Endobronchiální infiltrace bronchu B1 vpravo. Odebrán materiál k histologii."]},
                ("MUDr. Ondřej Kraus", "pneumolog"))),
            ("05_Histologie.docx", report(KOC_PAT, "Bioptický nález", p["name"], p["rc"], "Datum vydání: 14. 10. 2026", {
                "MIKROSKOPICKY": ["Bronchiální biopsie: infiltrace adenokarcinomem acinárního typu. "
                                  "EBUS 10R: metastáza adenokarcinomu v lymfatické uzlině. EBUS 4R a 7: bez nádoru."],
                "ZÁVĚR": ["Adenokarcinom plíce, acinární typ. Metastáza v uzlině 10R. "
                          "PD-L1 TPS 60 %, EGFR nemutovaný, ALK negativní. MKN-O 8140/3."]},
                ("MUDr. Radek Pokorný, Ph.D.", "patolog"))),
            ("06_MR_mozku.docx", report(KOC_RAD, "Výsledek vyšetření – MR mozku s k. l.", p["name"], p["rc"],
                                        "Datum vyšetření: 12. 10. 2026", {
                "ZÁVĚR": ["Mozek bez ložiskových změn, bez metastáz."]}, ("MUDr. Petra Malá", "radiolog"))),
            ("07_Laborator.docx", report(KOC_LAB, "Výsledky laboratorního vyšetření", p["name"], p["rc"],
                                         "Odběr: 2. 10. 2026 7:30", {
                "KREVNÍ OBRAZ": [[LAB_HEAD, ["Hemoglobin", "138", "g/l", "135–175"],
                                  ["Leukocyty", "8,9", "10⁹/l", "4,0–10,0"], ["Trombocyty", "312", "10⁹/l", "150–400"]]],
                "BIOCHEMIE": [[LAB_HEAD, ["Kreatinin", "88", "µmol/l", "64–104"], ["ALT", "0,41", "µkat/l", "0,10–0,78"],
                               ["Vápník", "2,41", "mmol/l", "2,15–2,55"]]]},
                ("Ing. Martina Šebestová", "odborný pracovník laboratoře"))),
            ("08_Funkcni_vysetreni_plic.docx", report(
                ["KOC Nemocnice Lipová (fiktivní)", "Pneumologická ambulance", "Nemocniční 15, 140 00 Praha 4"],
                "Ambulantní zpráva – funkční vyšetření plic", p["name"], p["rc"], "Datum vyšetření: 7. 10. 2026", {
                    "POPIS": ["Spirometrie: FEV1 2,41 l (78 % n. h.), FVC 3,52 l (86 % n. h.), FEV1/FVC 0,68. "
                              "Difuzní kapacita DLCO 74 % n. h. Spiroergometrie: VO2max 19,2 ml/kg/min."],
                    "ZÁVĚR": ["Lehká obstrukční ventilační porucha, mírně snížená DLCO. Předpokládané pooperační FEV1 "
                              "a DLCO nad 60 %, VO2max > 15 ml/kg/min – schopen lobektomie."]},
                ("MUDr. Ondřej Kraus", "pneumolog"))),
        ],
    },
    {
        "code": "breast", "name": "Eva Malá", "rc": "715608/1236", "who": "Pacientka", "referral": "1. 10. 2026",
        "docs": lambda p: [
            referral(p, "C50.4", ["Hmatná rezistence v levém prsu, zjištěno při samovyšetření.",
                                  "Osobní anamnéza: hypertenze, jinak zdráva. Rodinná anamnéza: matka karcinom prsu "
                                  "v 52 letech, jiné nádory v rodině neguje.",
                                  "Gynekologická anamnéza: postmenopauzální, poslední menstruace 2021."],
                     "Mamografie a UZ prsů 24. 9. 2026 (KOC Nemocnice Lipová) – BI-RADS 5 vlevo."),
            ("02_Mamografie.docx", report(KOC_RAD, "Výsledek vyšetření – mamografie a UZ prsů a axil", p["name"], p["rc"],
                                          "Datum vyšetření: 24. 9. 2026", {
                "POPIS": ["V horním zevním kvadrantu levého prsu spikulované ložisko 24 × 20 mm. "
                          "V levé axile zvětšená uzlina 15 mm se ztluštělým kortexem, suspektní. "
                          "Pravý prs bez ložiskových změn."],
                "ZÁVĚR": ["Ložisko levého prsu BI-RADS 5, suspektní uzlina v levé axile."]},
                ("MUDr. Petra Malá", "radiolog"), "Pacientka")),
            ("03_Histologie.docx", report(KOC_PAT, "Bioptický nález", p["name"], p["rc"], "Datum vydání: 3. 10. 2026", {
                "ZÁVĚR": ["Invazivní karcinom levého prsu NST, G2. ER 90 %, PR 70 %, HER2 negativní, Ki-67 25 %. "
                          "Uzlina levé axily: metastáza karcinomu. MKN-O 8500/3."]},
                ("MUDr. Radek Pokorný, Ph.D.", "patolog"), "Pacientka")),
            ("04_CT_staging.docx", report(KOC_RAD, "Výsledek vyšetření – CT hrudníku a břicha s k. l.", p["name"], p["rc"],
                                          "Datum vyšetření: 9. 10. 2026", {
                "ZÁVĚR": ["Plíce bez ložiskových změn. Játra bez ložiskových změn. Skelet bez ložisek. "
                          "Bez známek vzdálených metastáz."]}, ("MUDr. Petra Malá", "radiolog"), "Pacientka")),
            ("05_Laborator.docx", report(KOC_LAB, "Výsledky laboratorního vyšetření", p["name"], p["rc"],
                                         "Odběr: 2. 10. 2026 7:30", {
                "KREVNÍ OBRAZ": [[LAB_HEAD, ["Hemoglobin", "131", "g/l", "120–160"],
                                  ["Leukocyty", "6,4", "10⁹/l", "4,0–10,0"], ["Trombocyty", "265", "10⁹/l", "150–400"]]],
                "BIOCHEMIE": [[LAB_HEAD, ["Kreatinin", "71", "µmol/l", "44–80"], ["ALT", "0,32", "µkat/l", "0,10–0,55"],
                               ["ALP", "1,12", "µkat/l", "0,58–1,75"], ["Vápník", "2,36", "mmol/l", "2,15–2,55"]]]},
                ("Ing. Martina Šebestová", "odborný pracovník laboratoře"), "Pacientka")),
        ],
    },
]


def pathway_id(code: str) -> int:
    with session_scope() as s:
        return s.scalars(select(Pathway).where(Pathway.code == code)).one().id


def existing(rc: str) -> int | None:
    with session_scope() as s:
        c = s.scalars(select(Case).where(Case.patient_uuid == patient_uuid(rc))).first()
        return c.id if c else None


def walk(case_id: int) -> str:
    """Complete the manual diagnostic steps and decisions until the case is ready for the MDT."""
    for _ in range(8):
        d = svc.case_detail(case_id)
        if d["pre_mdt_complete"]:
            break
        todo = [n for n in d["pathway_graph"]["nodes"]
                if n.get("status") == "available" and n.get("phase") in ("intake", "diagnostics")]
        acted = False
        for n in todo:
            if n.get("completion") == "manual":
                svc.complete_step(case_id, n["id"], COORD)
                acted = True
            elif n.get("type") == "decision" and n.get("branches"):
                svc.choose_decision(case_id, n["id"], [n["branches"][0]["key"]], COORD)
                acted = True
        if not acted:
            break
    d = svc.case_detail(case_id)
    if not d["pre_mdt_complete"]:
        left = [n["id"] for n in d["pathway_graph"]["nodes"] if n.get("phase") in ("intake", "diagnostics")
                and n.get("status") not in ("done", "skipped") and n.get("blocking") is not False]
        print(f"   still open: {left}")
    return d["state"]


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--bodymap", action="store_true", help="also run the LLM body map")
    args = ap.parse_args()
    seed()  # centres and pathways, so this also works on a fresh install without the Rakathon sample data
    ids = []
    for p in PATIENTS:
        case_id = existing(p["rc"])
        if case_id:
            print(f"{p['name']} ({p['code']}) already there → case #{case_id}")
        else:
            files = p["docs"](p)
            case_id = svc.create_referral(KOC, files, actor=GP, sync=True, pathway_id=pathway_id(p["code"]),
                                          patient_name=p["name"], rc=p["rc"], referrer={"name": GP, "practice": PRACTICE})
            svc.accept_case(case_id, COORD)
            svc.set_mdt_date(case_id, MDT_DATE, COORD)
            print(f"{p['name']} ({p['code']}) referred with {len(files)} reports → case #{case_id}")
        print(f"   state: {walk(case_id)}")
        ids.append(case_id)
    if args.bodymap:
        from backend.reports import bodymap
        for case_id in ids:
            r = bodymap.compute(case_id)
            print(f"body map #{case_id}: " + ", ".join(f"{s['label_en']}={s['status']}" for s in r["sites"]))
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    raise SystemExit(main())

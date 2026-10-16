"""MDT documents.

* build_report(): the "Zápis k jednání MDT" (structure of MDT_zapis.docx) from the patient graph, using only
  documents dated on/before the MDT date (later reports, e.g. pain/psychology after the MDT, are excluded).
* parse_conclusion(): read an uploaded MDT conclusion (like MDT zaver.docx) into structured parts and suggest the
  pathway branch (rules first, MedGemma only as fallback; a human confirms the branch).
"""
from __future__ import annotations

import re
from datetime import date
from pathlib import Path

import docx
from docx.shared import Pt

from backend.extract import medgemma
from backend.extract.schemas import TreatmentChoice
from backend.graph.store import graph
from backend.ingest.reader import DocText, fold
from backend.reports import handover as hv


def build_report(case, koc, patient_uuid: str, mdt_date: date, entries: list[dict]) -> dict:
    h = hv.build(patient_uuid, mdt_date, entries, case.referring_doctor)
    f = {fld["key"]: fld for s in h["sections"] for fld in s["fields"]}
    patient = graph.patient(patient_uuid)
    facts = graph.facts(patient_uuid, mdt_date)

    def val(key, fallback="— chybí —"):
        return f[key]["text"] if key in f and f[key]["text"] else fallback

    dx = facts.get("diagnosis", [])
    confirmed = next((d["props"] for d in dx if d["props"].get("status") == "confirmed"), None)
    suspected = next((d["props"] for d in dx if d["props"].get("icd10")), None)
    code = (suspected or {}).get("icd10", "")
    if confirmed:
        dx_name = confirmed["name"]
        dx_note = "Kód MKN-10 převzat z žádosti PL – ověřit na MDT." if code else "Kód MKN-10 chybí."
    elif suspected:
        dx_name, dx_note = f"Suspektní: {suspected['name']}", "Histologie zatím není k dispozici."
    else:
        dx_name, dx_note = "— chybí —", ""

    rows = []
    for e in sorted(facts.get("exam", []), key=lambda e: e["props"].get("date") or date.min):
        p = e["props"]
        result = p.get("conclusion") or ""
        rows.append({"date": p.get("date"), "exam": f"{p.get('label')} ({p.get('place')})", "result": result,
                     "doc_id": e["sources"][0]["doc_id"] if e["sources"] else None})
    excluded = [d for d in graph.documents(patient_uuid) if d["d"].get("date") and d["d"]["date"] > mdt_date]

    rd = case.referring_doctor or {}
    report = {
        "mdt_date": mdt_date,
        "koc": koc.name,
        "patient": {"name": patient.get("patient_name", ""), "insurer": patient.get("insurer", ""),
                    "phone": patient.get("phone", "— chybí —"), "email": val("patient_email"),
                    "address": patient.get("address", ""), "weight_height": val("weight_height")},
        "referring_doctor": {"name": rd.get("name", ""), "phone": rd.get("phone") or "— chybí —",
                             "email": val("referring_doctor_email")},
        "members": koc.mdt_members or [],
        "diagnosis": {"name": dx_name, "icd10": code, "note": dx_note},
        "comorbidities": [i["text"] for i in f["comorbidities"]["items"]] if "comorbidities" in f else [],
        "results": rows,
        "performance_status": val("performance_status"),
        "symptoms": val("symptoms", ""),
        "medications": val("medications", ""),
        "mdt_question": val("mdt_question", ""),
        "notes": [i["text"] for i in f["notes"]["items"] if not i["text"].startswith(("RA:", "PA:", "Abusus:"))],
        "missing": h["missing"],
        "excluded_documents": [{"title": d["d"].get("title"), "date": d["d"].get("date")} for d in excluded],
        "conclusion": "",
        "recommendation": "",
        "treatments": [],
        "present": [],
        "final": False,
    }
    # After the meeting the same record becomes the final MDT conclusion (sections 6/7 filled).
    mc = case.mdt_conclusion or {}
    if mc.get("conclusion"):
        decision = next((n for n in case.pathway_version.graph_json["nodes"] if n["id"] == "decision"), {})
        labels = {b["key"]: b["label"] for b in decision.get("branches", [])}
        report.update({
            "final": True,
            "conclusion": mc["conclusion"],
            "recommendation": "\n".join(r for r in mc.get("recommendations", []) if r.strip()),
            "treatments": [labels.get(k, k) for k in (case.decisions or {}).get("decision", {}).get("keys", [])],
            "present": mc.get("members", []),
            "decided_by": mc.get("recorded_by", ""),
        })
        if mc.get("diagnosis"):
            report["diagnosis"]["mdt"] = mc["diagnosis"]
    return report


def _d(x) -> str:
    return f"{x.day}. {x.month}. {x.year}" if x else ""


def write_docx(report: dict, path: Path) -> Path:
    doc = docx.Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    final = report.get("final")
    doc.add_heading("Zápis z jednání multidisciplinárního týmu (MDT)" if final else
                    "Zápis k jednání multidisciplinárního týmu (MDT)", level=1)
    doc.add_paragraph(f"{report['koc']} · {'závěr MDT' if final else 'podklad pro MDT'} · vygenerováno aplikací Onkomost "
                      f"z dokumentace dostupné k {_d(report['mdt_date'])}")

    doc.add_heading("1  Základní údaje", level=2)
    p, r = report["patient"], report["referring_doctor"]
    t = doc.add_table(rows=0, cols=2)
    t.style = "Light Grid Accent 1"
    for k, v in [("Datum MDT", _d(report["mdt_date"])), ("Pacient", p["name"]), ("Pojišťovna", p["insurer"]),
                 ("Telefon, e-mail", f"{p['phone']}, {p['email']}"), ("Adresa", p["address"]),
                 ("Váha, výška", p["weight_height"]), ("Referující lékař", r["name"]),
                 ("Telefon, e-mail (lékař)", f"{r['phone']}, {r['email']}")]:
        row = t.add_row().cells
        row[0].text, row[1].text = k, v or "— chybí —"

    doc.add_heading("2  Členové MDT", level=2)
    present = set(report.get("present") or [])
    for m in report["members"] or [{"name": "— doplnit —", "role": ""}]:
        if final and present and m["name"] not in present:
            continue  # the final record lists who was actually present
        doc.add_paragraph(f"{m['name']} – {m['role']}" if m.get("role") else m["name"], style="List Bullet")

    doc.add_heading("3  Diagnóza", level=2)
    if report["diagnosis"].get("mdt"):
        doc.add_paragraph(report["diagnosis"]["mdt"])
    doc.add_paragraph(f"Název diagnózy: {report['diagnosis']['name']}")
    doc.add_paragraph(f"Kód diagnózy (MKN-10): {report['diagnosis']['icd10'] or '— chybí —'}")
    if report["diagnosis"]["note"]:
        doc.add_paragraph(report["diagnosis"]["note"]).italic = True

    doc.add_heading("4  Komorbidity", level=2)
    for c in report["comorbidities"] or ["— chybí —"]:
        doc.add_paragraph(c, style="List Bullet")

    doc.add_heading("5  Výsledky stávajících vyšetření", level=2)
    t = doc.add_table(rows=1, cols=3)
    t.style = "Light Grid Accent 1"
    t.rows[0].cells[0].text, t.rows[0].cells[1].text, t.rows[0].cells[2].text = "Datum", "Vyšetření", "Výsledek"
    for row in report["results"]:
        c = t.add_row().cells
        c[0].text, c[1].text, c[2].text = _d(row["date"]), row["exam"], row["result"]
    doc.add_paragraph(f"Stav výkonnosti: {report['performance_status']}")
    if report["mdt_question"]:
        doc.add_paragraph(f"Otázka referujícího lékaře: {report['mdt_question']}")
    if report["notes"]:
        doc.add_paragraph("Důležitá sdělení:")
        for n in report["notes"]:
            doc.add_paragraph(n, style="List Bullet")
    if report["missing"]:
        doc.add_paragraph("Chybějící údaje: " + ", ".join(m["label"] for m in report["missing"]))
    if report["excluded_documents"]:
        doc.add_paragraph("Nezahrnuto (vzniklo po datu MDT): " + ", ".join(
            f"{d['title']} ({_d(d['date'])})" for d in report["excluded_documents"]))

    doc.add_heading("6  Závěr MDT", level=2)
    doc.add_paragraph(report.get("conclusion") or "\n\n")
    if report.get("treatments"):
        doc.add_paragraph(f"Rozhodnutí o léčbě: {' + '.join(report['treatments'])}").runs[0].bold = True
    doc.add_heading("7  Doporučení MDT", level=2)
    recs = [r for r in (report.get("recommendation") or "").splitlines() if r.strip()]
    if recs:
        for r in recs:
            doc.add_paragraph(r)
    else:
        doc.add_paragraph("\n\n")
    if final:
        doc.add_paragraph(f"Zapsal/a: {report.get('decided_by') or '—'} · {_d(report['mdt_date'])}")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


# Treatment keys → patterns, applied only to the MDT's *primary* recommendation (the "doporučuje …" sentence
# up to "a následné / znovuposouzení" and the first recommendation line), so that a later re-assessment
# ("…znovuposouzením možnosti chirurgické resekce") does not count as a decision for surgery.
TREATMENT_RULES = [
    ("chemotherapy", r"chemoterapi|chemoradio|\bcht\b"),
    ("radiotherapy", r"radioterapi|ozareni|chemoradio|stereotakt"),
    ("surgery", r"resekc|operac|chirurgick|pankreatoduodenektom|lobektom|mastektom"),
    ("palliative", r"paliativ|symptomatick|podpurna lecba"),
    ("active_monitoring", r"aktivni sledovani|observac|watch"),
]


def primary_recommendation(conclusion: str, recommendations: list[str]) -> str:
    parts = []
    for sent in re.split(r"(?<=\.)\s+", conclusion or ""):
        f = fold(sent)
        if "doporucuje" in f:
            parts.append(re.split(r"\sa\snasledn|znovuposouz|po\s+neoadjuvan", f)[0])
    if recommendations:
        parts.append(re.split(r"[:–]", fold(recommendations[0]))[0])
    return " ".join(parts)


def suggest_treatments(conclusion: str, recommendations: list[str], allowed: set[str]) -> tuple[list[str], str | None]:
    text = primary_recommendation(conclusion, recommendations)
    keys = [k for k, pat in TREATMENT_RULES if k in allowed and re.search(pat, text)]
    if keys:
        return keys, "rule"
    if conclusion:
        try:
            r = medgemma.extract(TreatmentChoice, "Which treatments did the MDT recommend now? Allowed keys: "
                                 + ", ".join(sorted(allowed)), conclusion)
            keys = [k for k in r.treatments if k in allowed]
            if keys:
                return keys, "medgemma"
        except medgemma.LLMUnavailable:
            pass
    return [], None


def parse_conclusion(doc: DocText, branches: list[dict]) -> dict:
    conclusion = doc.section("ZÁVĚR MDT")
    recommendation = doc.section("DOPORUČENÍ MDT")
    diagnosis = doc.section("DIAGNÓZA")
    date_m = re.search(r"Datum MDT:\s*(\d{1,2}\.\s*\d{1,2}\.\s*\d{4})", doc.full_text)
    recs = [r.strip() for r in recommendation.splitlines() if r.strip()]
    keys, method = suggest_treatments(conclusion, recs, {b["key"] for b in branches})
    members = [l.split("–")[0].strip() for l in doc.section("ČLENOVÉ MDT").splitlines() if l.strip()]
    next_appt = next((r for r in recs if re.search(r"termín|\d{1,2}\.\s*\d{1,2}\.\s*\d{4}", r)), None)
    return {"date": date_m.group(1) if date_m else None, "diagnosis": diagnosis, "conclusion": conclusion,
            "recommendations": recs, "next_appointment": next_appt, "suggested_treatments": keys, "method": method,
            "members": members}

"""MDT case summary – what the team needs on one sheet to discuss and decide: the patient, the cancer, key findings,
lab trends, the patient's journey, the pathway status and open issues.

Assembled from the patient graph like the handover (no LLM), only from documents dated on/before the MDT date.
build() returns a list of blocks; the same blocks are written as an editable Word document (tables + charts) and as
HTML / plain text for the clipboard, so the download and the copy always say the same.

Block kinds: ("h", text) · ("p", text) · ("kv", [(label, value)]) · ("list", items) · ("img", png bytes, caption)
             · ("table", header, rows[, {"widths": [cm…], "shade": {row: (column, hex)}}])
"""
from __future__ import annotations

import base64
import html
import io
import re
from datetime import date

import matplotlib

matplotlib.use("Agg")
import matplotlib.dates as mdates  # noqa: E402
import matplotlib.pyplot as plt  # noqa: E402
import docx  # noqa: E402
from docx.enum.section import WD_ORIENT  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

from backend.graph.store import graph  # noqa: E402
from backend.ingest.identity import age_on  # noqa: E402
from backend.reports import handover as hv  # noqa: E402

L = {
    "cs": {
        "title": "Souhrn pacienta pro MDT", "generated": "Vygenerováno", "cutoff": "Zahrnuty dokumenty do",
        "patient": "Pacient", "name": "Jméno", "age": "Věk", "years": "let", "born": "nar.", "sex": "Pohlaví",
        "F": "žena", "M": "muž", "rc": "Rodné číslo", "insurer": "Pojišťovna", "phone": "Telefon",
        "contact": "Kontaktní osoba", "referrer": "Referující lékař", "koc": "Centrum (KOC)", "mdt_date": "Datum MDT",
        "days": "Od doporučení do MDT", "days_unit": "dní",
        "cancer": "Nádorové onemocnění", "type": "Typ nádoru (cesta)", "dx": "Diagnóza", "dx_confirmed": "histologicky ověřeno",
        "dx_suspected": "suspektní", "icd": "MKN-10", "morph": "Morfologie (MKN-O)", "grade": "Grade",
        "pni": "Perineurální invaze", "yes": "ano", "no": "ne", "imaging": "Hlavní zobrazovací nález",
        "markers": "Nádorové markery", "question": "Otázka / požadavek na MDT", "histology": "Histologie",
        "status": "Klinický stav", "ecog": "Výkonnostní stav (ECOG)", "not_stated": "není uveden v dokumentaci – doplnit",
        "weight": "Váha a výška", "symptoms": "Symptomy", "comorb": "Komorbidity", "meds": "Medikace", "allergy": "Alergie",
        "exams": "Provedená vyšetření", "date": "Datum", "exam": "Vyšetření", "place": "Pracoviště",
        "result": "Závěr", "labs": "Laboratoř – hodnoty mimo normu a nádorové markery", "analyte": "Parametr",
        "latest": "Poslední hodnota", "previous": "Předchozí", "ref": "Ref. rozmezí", "flag": "", "high": "↑ vysoká",
        "low": "↓ nízká", "lab_chart": "Vývoj laboratorních hodnot (šedý pás = referenční rozmezí, červeně mimo normu)",
        "journey": "Cesta pacienta", "journey_chart": "Dokumenty v čase podle typu; svislé čáry: doporučení do KOC a MDT",
        "referral_line": "doporučení", "pathway": "Stav onkologické cesty", "step": "Krok", "state": "Stav",
        "optional": "volitelné", "have": "doloženo", "missing": "chybí", "conflict": "nesoulad – ověřit",
        "skipped": "přeskočeno", "conditional": "podle rozhodnutí", "done": "hotovo", "open": "otevřené",
        "issues": "Před rozhodnutím ověřit", "no_issues": "Bez otevřených problémů.", "notes": "Důležitá sdělení",
        "discussion": "Diskuse a rozhodnutí MDT", "treatment": "Navržená léčba", "conclusion": "Závěr MDT",
        "recs": "Doporučení", "remarks": "Poznámky z diskuse", "decided": "Rozhodnutí zapsáno",
        "cat": {"referral": "Žádost a PL", "imaging": "Zobrazovací", "lab": "Laboratoř", "pathology": "Biopsie a histologie",
                "specialist": "Odborné zprávy", "mdt": "MDT", "other": "Ostatní"},
        "doc_conflict": "Nesoulad v dokumentu", "doc_missing": "Chybí dokument", "field_missing": "Chybí údaj",
    },
    "en": {
        "title": "MDT patient summary", "generated": "Generated", "cutoff": "Documents included up to",
        "patient": "Patient", "name": "Name", "age": "Age", "years": "years", "born": "born", "sex": "Sex",
        "F": "female", "M": "male", "rc": "Birth number (RČ)", "insurer": "Health insurer", "phone": "Phone",
        "contact": "Contact person", "referrer": "Referring doctor", "koc": "Cancer centre (KOC)", "mdt_date": "MDT date",
        "days": "Referral to MDT", "days_unit": "days",
        "cancer": "Cancer", "type": "Cancer type (pathway)", "dx": "Diagnosis", "dx_confirmed": "histologically confirmed",
        "dx_suspected": "suspected", "icd": "ICD-10", "morph": "Morphology (ICD-O)", "grade": "Grade",
        "pni": "Perineural invasion", "yes": "yes", "no": "no", "imaging": "Key imaging finding",
        "markers": "Tumour markers", "question": "Question / request to the MDT", "histology": "Histology",
        "status": "Clinical status", "ecog": "Performance status (ECOG)", "not_stated": "not stated in the documents – add",
        "weight": "Weight and height", "symptoms": "Symptoms", "comorb": "Comorbidities", "meds": "Medication",
        "allergy": "Allergies", "exams": "Examinations", "date": "Date", "exam": "Examination", "place": "Facility",
        "result": "Conclusion", "labs": "Labs – values outside the range and tumour markers", "analyte": "Analyte",
        "latest": "Latest value", "previous": "Previous", "ref": "Reference", "flag": "", "high": "↑ high", "low": "↓ low",
        "lab_chart": "Lab values over time (grey band = reference range, red = outside the range)",
        "journey": "Patient journey", "journey_chart": "Documents over time by type; vertical lines: referral to the KOC and the MDT",
        "referral_line": "referral", "pathway": "Cancer pathway status", "step": "Step", "state": "Status",
        "optional": "optional", "have": "documented", "missing": "missing", "conflict": "conflict – check",
        "skipped": "skipped", "conditional": "depends on a decision", "done": "done", "open": "open",
        "issues": "Check before deciding", "no_issues": "No open issues.", "notes": "Important notes",
        "discussion": "MDT discussion and decision", "treatment": "Proposed treatment", "conclusion": "MDT conclusion",
        "recs": "Recommendations", "remarks": "Discussion notes", "decided": "Decision recorded",
        "cat": {"referral": "Referral & GP", "imaging": "Imaging", "lab": "Labs", "pathology": "Biopsy & histology",
                "specialist": "Specialist reports", "mdt": "MDT", "other": "Other"},
        "doc_conflict": "Conflict in document", "doc_missing": "Missing document", "field_missing": "Missing data",
    },
}
CATEGORIES = ["referral", "imaging", "lab", "pathology", "specialist", "mdt", "other"]
CAT_COLOR = {"referral": "#9b9a97", "imaging": "#337ea9", "lab": "#6f7fd1", "pathology": "#d44c47",
             "specialist": "#d9730d", "mdt": "#0f7b6c", "other": "#c7c6c3"}
IMAGING = ("ct", "mri", "pet_ct", "ultrasound", "xray", "mammography")
STATUS_OK, STATUS_WARN, STATUS_BAD = "E6F4EA", "FDF3D8", "FBE4E4"


def _d(x) -> str:
    return f"{x.day}. {x.month}. {x.year}" if x else "—"


def category(doc_type: str) -> str:
    """Same rows as the journey timeline in the UI (frontend/src/facts.ts docCategory)."""
    if doc_type in ("referral", "gp_summary"):
        return "referral"
    if doc_type in ("imaging", "lab", "mdt_conclusion"):
        return {"mdt_conclusion": "mdt"}.get(doc_type, doc_type)
    if doc_type in ("procedure", "histology"):
        return "pathology"
    return "specialist" if doc_type.endswith("_report") else "other"


def _flag(t: dict, flag: str | None) -> str:
    """Lab flag "H" / "L" / "++" → "↑ high" / "↓ low"; nothing when the value is in range."""
    if not flag:
        return ""
    return t["high"] if re.search(r"H|\+", flag) else t["low"]


def _num(s: str) -> float | None:
    try:
        return float(s.replace(" ", "").replace(",", "."))
    except ValueError:
        return None


def parse_ref(ref: str) -> tuple[float | None, float | None]:
    """"0,58–1,75" → (0.58, 1.75); "< 35" → (None, 35); "> 1,0" → (1.0, None)."""
    r = (ref or "").strip()
    if m := re.match(r"^([\d.,]+)\s*[–-]\s*([\d.,]+)", r):
        return _num(m.group(1)), _num(m.group(2))
    if m := re.match(r"^[<≤]\s*([\d.,]+)", r):
        return None, _num(m.group(1))
    if m := re.match(r"^[>≥]\s*([\d.,]+)", r):
        return _num(m.group(1)), None
    return None, None


def lab_series(labs: list[dict]) -> list[dict]:
    """One series per analyte + unit (blood and urine leukocytes stay apart), oldest value first."""
    by: dict[tuple, dict] = {}
    for l in labs:
        p = l["props"]
        if not p.get("analyte") or not p.get("date"):
            continue
        s = by.setdefault((p["analyte"], p.get("unit") or ""), {
            "analyte": p["analyte"], "unit": p.get("unit") or "", "ref": p.get("ref") or "", "marker": bool(p.get("is_tumor_marker")),
            "pts": []})
        if not any(x["date"] == p["date"] and x["text"] == p.get("value") for x in s["pts"]):
            s["pts"].append({"date": p["date"], "value": p.get("value_num"), "text": str(p.get("value") or ""), "flag": p.get("flag")})
    for s in by.values():
        s["pts"].sort(key=lambda x: x["date"])
        s["lo"], s["hi"] = parse_ref(s["ref"])
    return list(by.values())


# --- charts ---------------------------------------------------------------------------------------------------------

def _png(fig) -> bytes:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=160, bbox_inches="tight")
    plt.close(fig)
    return buf.getvalue()


def _short(title: str) -> str:
    t = re.sub(r"^(Výsledek vyšetření|Ambulantní zpráva|Protokol o výkonu|Výsledky)\s*[–-]?\s*", "", title or "").strip()
    return t[:34] + ("…" if len(t) > 34 else "")


def timeline_chart(docs: list[dict], referral: date | None, mdt: date | None, lang: str) -> bytes | None:
    docs = [d for d in docs if d.get("date")]
    if not docs:
        return None
    t = L[lang]
    rows = [c for c in CATEGORIES if any(category(d.get("type") or "") == c for d in docs)]
    fig, ax = plt.subplots(figsize=(9.6, 0.55 * len(rows) + 1.1))
    for i, c in enumerate(rows):
        ds = sorted((d for d in docs if category(d.get("type") or "") == c), key=lambda d: d["date"])
        ax.scatter([d["date"] for d in ds], [i] * len(ds), s=46, color=CAT_COLOR[c], zorder=3)
        for k, d in enumerate(ds):  # alternate the labels above / below so neighbours do not overlap
            ax.annotate(_short(d.get("title") or ""), (d["date"], i), xytext=(4, 7 if k % 2 == 0 else -11),
                        textcoords="offset points", fontsize=6.3, color="#37352f")
    for when, label, style in ((referral, t["referral_line"], ":"), (mdt, "MDT", "--")):
        if when:
            ax.axvline(when, color="#0f7b6c" if label == "MDT" else "#787774", linestyle=style, linewidth=1.1, zorder=2)
            ax.annotate(label, (when, len(rows) - 0.45), xytext=(3, 0), textcoords="offset points", fontsize=7,
                        color="#0f7b6c" if label == "MDT" else "#787774", fontweight="bold")
    ax.set_yticks(range(len(rows)), [t["cat"][c] for c in rows], fontsize=8)
    ax.set_ylim(-0.7, len(rows) - 0.2)
    ax.invert_yaxis()
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%d. %m."))
    ax.tick_params(axis="x", labelsize=7.5)
    ax.grid(axis="x", color="#e9e9e7", linewidth=0.8)
    for side in ("top", "right", "left"):
        ax.spines[side].set_visible(False)
    return _png(fig)


def lab_chart(series: list[dict], lang: str) -> bytes | None:
    """Small multiples on one shared date axis: tumour markers first, then repeated values that left the range."""
    numeric = [s for s in series if any(p["value"] is not None for p in s["pts"])]
    picked = [s for s in numeric if s["marker"]] + sorted(
        (s for s in numeric if not s["marker"] and len(s["pts"]) > 1 and any(p["flag"] for p in s["pts"])),
        key=lambda s: -sum(1 for p in s["pts"] if p["flag"]))
    picked = picked[:9]
    if not picked:
        return None
    dates = [p["date"] for s in picked for p in s["pts"]]
    lo_d, hi_d = min(dates), max(dates)
    pad = max((hi_d - lo_d).days * 0.08, 3)
    cols = 3
    rows = (len(picked) + cols - 1) // cols
    fig, axes = plt.subplots(rows, cols, figsize=(9.6, 2.05 * rows), squeeze=False)
    for ax, s in zip(axes.flat, picked):
        pts = [p for p in s["pts"] if p["value"] is not None]
        xs, ys = [p["date"] for p in pts], [p["value"] for p in pts]
        if s["lo"] is not None or s["hi"] is not None:
            top = max(ys + [s["hi"] or 0]) * 1.15
            ax.axhspan(s["lo"] or 0, s["hi"] if s["hi"] is not None else top, color="#e9e9e7", zorder=0)
        ax.plot(xs, ys, color="#37352f", linewidth=1.2, zorder=2)
        ax.scatter(xs, ys, s=26, zorder=3, color=["#d44c47" if p["flag"] else "#37352f" for p in pts])
        for x, y, p in zip(xs, ys, pts):
            ax.annotate(p["text"], (x, y), xytext=(0, 5), textcoords="offset points", ha="center", fontsize=6.8)
        ax.set_title(f"{s['analyte']} ({s['unit']})" if s["unit"] else s["analyte"], fontsize=8, loc="left",
                     color="#d44c47" if s["marker"] else "#37352f")
        ax.set_xlim(matplotlib.dates.date2num(lo_d) - pad, matplotlib.dates.date2num(hi_d) + pad)
        ax.xaxis.set_major_formatter(mdates.DateFormatter("%d. %m."))
        ax.xaxis.set_major_locator(mdates.AutoDateLocator(maxticks=4))
        ax.tick_params(labelsize=6.5)
        ax.margins(y=0.25)
        if min(ys) >= 0:
            ax.set_ylim(bottom=0)
        for side in ("top", "right"):
            ax.spines[side].set_visible(False)
    for ax in list(axes.flat)[len(picked):]:
        ax.axis("off")
    fig.tight_layout(h_pad=1.4)
    return _png(fig)


# --- content --------------------------------------------------------------------------------------------------------

def build(case, cutoff: date, entries: list[dict], result: dict, generated_on: date, lang: str = "cs") -> dict:
    lang = lang if lang in L else "cs"
    t = L[lang]
    p_uuid = case.patient_uuid
    h = hv.build(p_uuid, cutoff, entries, case.referring_doctor)
    fld = {f["key"]: f for s in h["sections"] for f in s["fields"]}
    text = lambda key: fld[key]["text"] if key in fld and fld[key]["text"] else "—"  # noqa: E731
    facts = graph.facts(p_uuid, cutoff)
    patient = graph.patient(p_uuid)
    pathway = case.pathway_version.pathway
    blocks: list[tuple] = []

    # 1 patient
    dob = patient.get("dob")
    if dob is not None and hasattr(dob, "to_native"):
        dob = dob.to_native()
    days = (case.mdt_date - case.created_on).days if case.mdt_date else None
    blocks += [("h", t["patient"]), ("kv", [
        (t["name"], patient.get("patient_name", "")),
        (t["age"], f"{age_on(dob, cutoff)} {t['years']} ({t['born']} {_d(dob)})" if dob else "—"),
        (t["sex"], t.get(patient.get("sex", ""), "—")), (t["rc"], patient.get("rc", "—")),
        (t["insurer"], patient.get("insurer", "—")), (t["phone"], patient.get("phone", "—")),
        (t["contact"], text("contact_person")), (t["referrer"], text("referring_doctor")),
        (t["koc"], case.koc.name), (t["mdt_date"], _d(case.mdt_date)),
        *([(t["days"], f"{days} {t['days_unit']}")] if days is not None else []),
    ])]

    # 2 cancer
    dx = facts.get("diagnosis", [])
    confirmed = next((d for d in dx if d["props"].get("status") == "confirmed"), None)
    suspected = next((d for d in dx if d["props"].get("icd10")), None)
    rows = [(t["type"], pathway.name if lang == "cs" else (pathway.name_en or pathway.name))]
    if confirmed:
        cp = confirmed["props"]
        when = _d(confirmed["sources"][0]["doc_date"]) if confirmed["sources"] else ""
        rows.append((t["dx"], f"{cp.get('name')} ({t['dx_confirmed']} {when})"))
        rows += [(k, v) for k, v in ((t["morph"], cp.get("morphology")), (t["grade"], cp.get("grade"))) if v]
        if "perineural_invasion" in cp:
            rows.append((t["pni"], t["yes"] if cp["perineural_invasion"] else t["no"]))
    elif suspected:
        rows.append((t["dx"], f"{suspected['props'].get('name')} ({t['dx_suspected']})"))
    if suspected:
        rows.append((t["icd"], suspected["props"].get("icd10")))
    exams = sorted(facts.get("exam", []), key=lambda e: e["props"].get("date") or date.min)
    imaging = [e for e in exams if e["props"].get("exam_type") in IMAGING]
    key_img = next((e for e in reversed(imaging) if e["props"].get("exam_type") in ("ct", "mri", "pet_ct")), None) \
        or (imaging[-1] if imaging else None)
    if key_img:
        ip = key_img["props"]
        rows.append((t["imaging"], f"{ip.get('label')} ({_d(ip.get('date'))}): {ip.get('conclusion')}"))
    series = lab_series(facts.get("lab", []))
    markers = [s for s in series if s["marker"]]
    if markers:
        rows.append((t["markers"], "; ".join(
            f"{s['analyte']} {s['pts'][-1]['text']} {s['unit']} (ref. {s['ref']}) {_flag(t, s['pts'][-1]['flag'])}".rstrip()
            + f" – {_d(s['pts'][-1]['date'])}" for s in markers)))
    if "histology" in fld and fld["histology"]["text"]:
        rows.append((t["histology"], fld["histology"]["text"]))
    rows.append((t["question"], text("mdt_question")))
    blocks += [("h", t["cancer"]), ("kv", rows)]

    # 3 clinical status
    ecog = fld.get("performance_status", {}).get("text") or t["not_stated"]
    blocks += [("h", t["status"]), ("kv", [
        (t["ecog"], ecog), (t["weight"], text("weight_height")), (t["symptoms"], text("symptoms").replace("\n", "; ")),
        (t["comorb"], text("comorbidities").replace("\n", "; ")), (t["meds"], text("medications").replace("\n", "; ")),
        (t["allergy"], text("allergies")),
    ])]

    # 4 examinations
    exam_rows = [(_d(e["props"].get("date")), e["props"].get("label") or "", e["props"].get("place") or "",
                  (e["props"].get("conclusion") or "").replace("\n", " ")) for e in exams if e["props"].get("exam_type") != "lab"]
    if exam_rows:
        blocks += [("h", t["exams"]), ("table", [t["date"], t["exam"], t["place"], t["result"]], exam_rows,
                                       {"widths": [2.3, 4.0, 3.9, 7.2]})]

    # 5 labs
    # tumour markers + what is out of range in the latest tests (an old urine dipstick is not what the MDT decides on)
    latest = max((s["pts"][-1]["date"] for s in series), default=None)
    lab_rows = []
    for s in sorted(series, key=lambda s: (not s["marker"], s["analyte"])):
        last = s["pts"][-1]
        if not (s["marker"] or (last["flag"] and last["date"] == latest)):
            continue
        prev = s["pts"][-2] if len(s["pts"]) > 1 else None
        lab_rows.append((s["analyte"], f"{last['text']} {s['unit']} ({_d(last['date'])})".strip(),
                         f"{prev['text']} ({_d(prev['date'])})" if prev else "—", s["ref"] or "—", _flag(t, last["flag"])))
    if lab_rows:
        blocks += [("h", t["labs"]), ("table", [t["analyte"], t["latest"], t["previous"], t["ref"], t["flag"]], lab_rows,
                                      {"widths": [3.8, 4.4, 3.6, 3.0, 2.6]})]
    chart = lab_chart(series, lang)
    if chart:
        blocks.append(("img", chart, t["lab_chart"]))

    # 6 journey
    docs = [d["d"] for d in graph.documents(p_uuid, cutoff) if d["d"].get("date")]
    # label each point by what it is ("Kardiologie – kontrolní vyšetření"), not by the letter's generic title
    names = {src["doc_uid"]: e["props"]["label"] for e in exams if e["props"].get("label") for src in e["sources"]}
    chart = timeline_chart([{**d, "title": names.get(d.get("uid"), d.get("title"))} for d in docs],
                           case.created_on, case.mdt_date, lang)
    if chart:
        blocks += [("h", t["journey"]), ("img", chart, t["journey_chart"])]

    # 7 pathway status and open issues
    g = case.pathway_version.graph_json
    steps, shade = [], {}
    for n in g["nodes"]:
        st = result["nodes"].get(n["id"], {})
        if n.get("phase") not in ("intake", "diagnostics", "supportive") or n.get("type") == "start":
            continue
        label = n.get("label") if lang == "cs" else (n.get("label_en") or n.get("label"))
        status = "skipped" if st.get("status") == "skipped" else st.get("doc_status") or (
            "done" if st.get("status") == "done" else "open")
        optional = n.get("blocking") is False
        # green = done, red = a required step is missing, amber = conflict; optional gaps stay neutral
        color = {"have": STATUS_OK, "done": STATUS_OK, "conflict": STATUS_WARN}.get(status) or (
            STATUS_BAD if status == "missing" and not optional else None)
        if color:
            shade[len(steps)] = (1, color)
        when = _d(st.get("done_on")) if st.get("done_on") else "—"
        steps.append((label + (f" ({t['optional']})" if optional else ""), t[status], when))
    blocks += [("h", t["pathway"]), ("table", [t["step"], t["state"], t["date"]], steps,
                                      {"widths": [9.6, 4.6, 3.2], "shade": shade})]
    issues = []
    for f in result["findings"]:
        if f["kind"] == "mismatch":  # the title already says "Conflict in document: <title>"
            issues.append(f"{f.get('title') if lang == 'cs' else f.get('title_en') or f.get('title')} – "
                          f"{f.get('detail') if lang == 'cs' else f.get('detail_en') or f.get('detail')}")
    for n in g["nodes"]:
        st = result["nodes"].get(n["id"], {})
        if st.get("doc_status") == "missing" and n.get("blocking") is not False:
            issues.append(f"{t['doc_missing']}: {n.get('label') if lang == 'cs' else n.get('label_en') or n.get('label')}")
    for m in h["missing"]:
        if m["required"]:  # clinical gaps (ECOG…), not the optional contact fields of the form
            issues.append(f"{t['field_missing']}: {m['label'] if lang == 'cs' else m['label_en']}")
    blocks += [("h", t["issues"]), ("list", issues or [t["no_issues"]])]

    # 8 notes
    notes = [i["text"] for i in fld.get("notes", {}).get("items", []) if not i["text"].startswith(("RA:", "PA:", "Abusus:"))]
    if notes:
        blocks += [("h", t["notes"]), ("list", notes)]

    # 9 discussion: filled once the MDT recorded its decision, otherwise empty rows to write in
    mc = case.mdt_conclusion or {}
    decision = next((n for n in g["nodes"] if n["id"] == "decision"), {})
    labels = {b["key"]: (b["label"] if lang == "cs" else b.get("label_en") or b["label"]) for b in decision.get("branches", [])}
    chosen = ", ".join(labels.get(k, k) for k in (case.decisions or {}).get("decision", {}).get("keys", []))
    blocks += [("h", t["discussion"]), ("kv", [
        (t["treatment"], chosen or ""), (t["conclusion"], mc.get("conclusion", "")),
        (t["recs"], "\n".join(mc.get("recommendations", []))), (t["remarks"], ""),
    ])]

    name = patient.get("patient_name", "")
    return {
        "title": t["title"], "patient": name, "lang": lang,
        "subtitle": f"{name} · {case.koc.name} · {t['mdt_date']} {_d(case.mdt_date)}",
        "meta": f"{t['generated']} {_d(generated_on)} · {t['cutoff']} {_d(cutoff)} · Onkomost",
        "blocks": blocks,
        "filename": f"MDT_souhrn_{re.sub(r'[^A-Za-z0-9]+', '_', _ascii(name)).strip('_')}_{(case.mdt_date or cutoff).isoformat()}.docx",
    }


def _ascii(s: str) -> str:
    import unicodedata
    return "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c))


# --- renderers ------------------------------------------------------------------------------------------------------

def _shade(cell, hex_color: str) -> None:
    tc = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc.append(shd)


def _cell(cell, text: str, bold: bool = False, size: float = 9) -> None:
    cell.text = ""
    par = cell.paragraphs[0]
    for i, line in enumerate(str(text or "").split("\n")):
        if i:
            par = cell.add_paragraph()
        run = par.add_run(line)
        run.bold, run.font.size = bold, Pt(size)
        par.paragraph_format.space_after = Pt(0)


def to_docx(s: dict) -> bytes:
    doc = docx.Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.6)
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Calibri", Pt(10)
    doc.add_heading(s["title"], level=1)
    doc.add_paragraph(s["subtitle"]).runs[0].bold = True
    meta = doc.add_paragraph(s["meta"])
    meta.runs[0].font.size, meta.runs[0].font.color.rgb = Pt(8.5), RGBColor(0x78, 0x77, 0x74)
    width = sec.page_width - sec.left_margin - sec.right_margin
    for b in s["blocks"]:
        kind = b[0]
        if kind == "h":
            doc.add_heading(b[1], level=2)
        elif kind == "p":
            doc.add_paragraph(b[1]).runs[0].font.size = Pt(9)
        elif kind == "list":
            for item in b[1]:
                doc.add_paragraph(item, style="List Bullet").runs[0].font.size = Pt(9)
        elif kind == "kv":
            table = doc.add_table(rows=0, cols=2)
            table.style = "Table Grid"
            for k, v in b[1]:
                row = table.add_row().cells
                _cell(row[0], k, bold=True)
                _cell(row[1], v or "")
                _shade(row[0], "F1F1EF")
                row[0].width, row[1].width = Cm(4.6), width - Cm(4.6)
        elif kind == "table":
            header, rows, opts = b[1], b[2], (b[3] if len(b) > 3 else {})
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            table.autofit = not opts.get("widths")
            for c, txt in zip(table.rows[0].cells, header):
                _cell(c, txt, bold=True)
                _shade(c, "E3EDF5")
            for i, r in enumerate(rows):
                cells = table.add_row().cells
                for c, txt in zip(cells, r):
                    _cell(c, txt, size=8.5)
                if i in opts.get("shade", {}):
                    col, color = opts["shade"][i]
                    _shade(cells[col], color)
            for row in table.rows if opts.get("widths") else []:  # Word reads the width of every cell, not the column
                for c, w in zip(row.cells, opts["widths"]):
                    c.width = Cm(w)
        elif kind == "img":
            doc.add_picture(io.BytesIO(b[1]), width=width)
            cap = doc.add_paragraph(b[2])
            cap.runs[0].italic, cap.runs[0].font.size = True, Pt(8)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def to_html(s: dict) -> str:
    """Self-contained HTML (inline styles, charts as data URIs) – pastes into Word, Outlook or Teams with its tables."""
    e = html.escape
    td = 'style="border:1px solid #d3d1cb;padding:4px 6px;vertical-align:top;font-size:10pt"'
    th = 'style="border:1px solid #d3d1cb;padding:4px 6px;background:#e3edf5;text-align:left;font-size:10pt"'
    out = [f'<div style="font-family:Calibri,Arial,sans-serif;font-size:10.5pt;color:#37352f">',
           f"<h1 style=\"font-size:16pt;margin:0 0 4px\">{e(s['title'])}</h1>",
           f"<p style=\"margin:0\"><b>{e(s['subtitle'])}</b></p>",
           f"<p style=\"margin:0 0 8px;color:#787774;font-size:8.5pt\">{e(s['meta'])}</p>"]
    br = lambda x: e(str(x or "")).replace("\n", "<br>")  # noqa: E731
    for b in s["blocks"]:
        kind = b[0]
        if kind == "h":
            out.append(f'<h2 style="font-size:12.5pt;margin:14px 0 6px">{e(b[1])}</h2>')
        elif kind == "p":
            out.append(f"<p>{br(b[1])}</p>")
        elif kind == "list":
            out.append("<ul>" + "".join(f"<li>{br(i)}</li>" for i in b[1]) + "</ul>")
        elif kind == "kv":
            out.append('<table style="border-collapse:collapse;width:100%">' + "".join(
                f'<tr><td {td.replace("vertical-align:top", "vertical-align:top;background:#f1f1ef;font-weight:bold;width:28%")}>'
                f"{e(k)}</td><td {td}>{br(v)}</td></tr>" for k, v in b[1]) + "</table>")
        elif kind == "table":
            shade = (b[3] if len(b) > 3 else {}).get("shade", {})
            cell = lambda i, j, c: (  # noqa: E731
                f'<td {td.replace("vertical-align:top", "vertical-align:top;background:#" + shade[i][1])}>{br(c)}</td>'
                if i in shade and shade[i][0] == j else f"<td {td}>{br(c)}</td>")
            out.append('<table style="border-collapse:collapse;width:100%"><tr>'
                       + "".join(f"<th {th}>{e(h)}</th>" for h in b[1]) + "</tr>"
                       + "".join("<tr>" + "".join(cell(i, j, c) for j, c in enumerate(r)) + "</tr>" for i, r in enumerate(b[2]))
                       + "</table>")
        elif kind == "img":
            out.append(f'<p><img alt="{e(b[2])}" style="max-width:100%" width="720" '
                       f'src="data:image/png;base64,{base64.b64encode(b[1]).decode()}"><br>'
                       f'<i style="font-size:8.5pt;color:#787774">{e(b[2])}</i></p>')
    out.append("</div>")
    return "".join(out)


def to_text(s: dict) -> str:
    out = [s["title"].upper(), s["subtitle"], s["meta"], ""]
    for b in s["blocks"]:
        kind = b[0]
        if kind == "h":
            out += ["", b[1].upper()]
        elif kind == "p":
            out.append(b[1])
        elif kind == "list":
            out += [f"- {i}" for i in b[1]]
        elif kind == "kv":
            out += [f"{k}: {v}" if v else f"{k}:" for k, v in b[1]]
        elif kind == "table":
            out.append(" | ".join(b[1]))
            out += [" | ".join(str(c).replace("\n", " ") for c in r) for r in b[2]]
        elif kind == "img":
            out.append(f"[{b[2]}]")
    return "\n".join(out)
